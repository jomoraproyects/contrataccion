from datetime import date, datetime, timedelta
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .models import Perfil, ProcesoSeleccion, SeguimientoEtapa
from .forms import IndicadoresFiltroForm
from .reportes import calcular_indicadores


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class IndicadoresPanelTests(TestCase):
    def setUp(self):
        self.creador = User.objects.create_user("creador-panel", password="temporal")
        self.creador.perfil.rol = Perfil.Rol.VINCULACION
        self.creador.perfil.save()
        self.gerente = User.objects.create_user("gerente-panel", password="temporal")
        self.gerente.perfil.rol = Perfil.Rol.GERENTE
        self.gerente.perfil.save()
        self.client.force_login(self.gerente)

    def candidato(self, numero, fecha=date(2026, 9, 8), vacante="Ventas"):
        return ProcesoSeleccion.objects.create(
            nombre=f"Persona {numero}", apellidos="Panel", cedula=f"99881{numero:03d}",
            celular="3001234567", fecha_llegada=fecha, vacante=vacante,
            creado_por=self.creador,
        )

    def test_graficas_y_exportacion_usan_los_mismos_filtros(self):
        self.candidato(1, vacante="Ventas")
        self.candidato(2, vacante="VENTAS")
        self.candidato(3, vacante="Compras")
        self.candidato(4, fecha=date(2025, 1, 1))
        filtros = {"fecha_desde": "2026-09-01", "fecha_hasta": "2026-09-30", "vacante": "vEnTaS", "etapa": "RRHH", "estado": "EN_CURSO"}
        respuesta = self.client.get(reverse("procesos:indicadores"), filtros)
        graficas = respuesta.context["indicadores"]["graficas"]
        self.assertEqual(graficas["total"], 2)
        self.assertEqual(sum(item["cantidad"] for item in graficas["estados"]), 2)
        self.assertEqual(graficas["vacantes"], [{"etiqueta": "VENTAS", "cantidad": 2}])
        excel = self.client.get(reverse("procesos:exportar_indicadores", args=["xlsx"]), filtros)
        hoja = load_workbook(BytesIO(excel.content), read_only=True)["Candidatos"]
        self.assertEqual(hoja.max_row, 3)
        self.assertContains(respuesta, "datos-indicadores")
        self.assertContains(respuesta, "js/indicadores.js?v=20261028.3")

    def test_serie_mensual_continua_y_limitada_a_doce_meses(self):
        self.candidato(1, fecha=date(2025, 1, 1))
        self.candidato(2, fecha=date(2026, 7, 1))
        self.candidato(3, fecha=date(2026, 9, 1))
        serie = calcular_indicadores(ProcesoSeleccion.objects.prefetch_related("seguimientos"))["graficas"]["tendencia"]
        self.assertEqual(len(serie), 12)
        self.assertEqual(serie[0]["etiqueta"], "10/2025")
        self.assertEqual(serie[-2], {"etiqueta": "08/2026", "cantidad": 0})
        self.assertEqual(serie[-1], {"etiqueta": "09/2026", "cantidad": 1})

    def test_cumplimiento_no_confunde_abiertos_con_finalizados(self):
        inicio = timezone.make_aware(datetime(2026, 9, 8, 7))
        for numero, horas in enumerate((1, 2, 9), start=1):
            seguimiento = self.candidato(numero).seguimientos.get()
            seguimiento.inicio = inicio
            seguimiento.fin = inicio + timedelta(hours=horas)
            seguimiento.fecha_limite = inicio + timedelta(hours=4)
            seguimiento.resultado = SeguimientoEtapa.Resultado.APROBADO
            seguimiento.save()
        pendiente = self.candidato(4).seguimientos.get()
        pendiente.fecha_limite = timezone.now() - timedelta(days=1)
        pendiente.save()
        ajuste = self.candidato(5).seguimientos.get()
        ajuste.fin = timezone.now()
        ajuste.resultado = SeguimientoEtapa.Resultado.AJUSTADO
        ajuste.save()
        resumen = calcular_indicadores(ProcesoSeleccion.objects.prefetch_related("seguimientos"))
        etapa = resumen["graficas"]["etapas"][0]
        self.assertEqual(etapa["promedio_horas"], 4)
        self.assertEqual(etapa["mediana_horas"], 2)
        self.assertEqual(etapa["a_tiempo"], 2)
        self.assertEqual(etapa["retrasos_cerrados"], 1)
        self.assertEqual(etapa["cumplimiento_porcentaje"], 67)
        self.assertEqual(resumen["vencidos"], 1)
        self.assertEqual(resumen["etapas_finalizadas"], 3)
        self.assertEqual(etapa["plazo_horas"], 18)

    def test_datos_del_grafico_escapan_texto_y_filtros_invalidos_vacian_reporte(self):
        self.candidato(1, vacante="</script><script>alert(1)</script>")
        respuesta = self.client.get(reverse("procesos:indicadores"), {"periodo": "mensual", "mes": "2026-09"})
        self.assertNotContains(respuesta, "</script><script>alert(1)</script>")
        invalida = self.client.get(reverse("procesos:indicadores"), {"fecha_desde": "2026-09-30", "fecha_hasta": "2026-09-01"})
        self.assertEqual(invalida.context["indicadores"]["graficas"]["total"], 0)
        self.assertEqual(invalida.context["indicadores"]["cumplimiento_porcentaje"], None)
        self.assertEqual(self.client.get(reverse("procesos:exportar_indicadores", args=["pdf"]), {"estado": "NO_EXISTE"}).status_code, 400)

    def test_periodos_resuelven_semana_iso_mes_bisiesto_y_anio_completo(self):
        casos = [
            ({"periodo": "semanal", "semana": "2026-W53"}, date(2026, 12, 28), date(2027, 1, 3)),
            ({"periodo": "mensual", "mes": "2024-02"}, date(2024, 2, 1), date(2024, 2, 29)),
            ({"periodo": "anual", "anio": "2025"}, date(2025, 1, 1), date(2025, 12, 31)),
        ]
        for entrada, inicio, fin in casos:
            with self.subTest(entrada=entrada):
                formulario = IndicadoresFiltroForm(entrada)
                self.assertTrue(formulario.is_valid(), formulario.errors)
                self.assertEqual(formulario.cleaned_data["fecha_desde"], inicio)
                self.assertEqual(formulario.cleaned_data["fecha_hasta"], fin)

    def test_periodos_invalidos_no_exportan(self):
        for filtros in (
            {"periodo": "semanal", "semana": "2026-W54"},
            {"periodo": "semanal", "semana": "9999-W52"},
            {"periodo": "mensual", "mes": "2026-13"},
            {"periodo": "anual", "anio": "9000"},
            {"periodo": "todo"},
        ):
            with self.subTest(filtros=filtros):
                self.assertFalse(IndicadoresFiltroForm(filtros).is_valid())
                respuesta = self.client.get(reverse("procesos:exportar_indicadores", args=["xlsx"]), filtros)
                self.assertEqual(respuesta.status_code, 400)

    def test_semana_cruza_anio_y_exporta_exactamente_los_candidatos_del_periodo(self):
        for indice, fecha in enumerate((date(2026, 12, 27), date(2026, 12, 28), date(2027, 1, 3), date(2027, 1, 4))):
            self.candidato(indice + 1, fecha=fecha)
        filtros = {"periodo": "semanal", "semana": "2026-W53"}
        respuesta = self.client.get(reverse("procesos:indicadores"), filtros)
        grafica = respuesta.context["indicadores"]["graficas"]
        self.assertEqual(grafica["total"], 2)
        self.assertEqual(len(grafica["tendencia"]), 7)
        self.assertEqual(grafica["tendencia"][0], {"etiqueta": "28/12", "cantidad": 1})
        self.assertEqual(grafica["tendencia"][-1], {"etiqueta": "03/01", "cantidad": 1})
        excel = self.client.get(reverse("procesos:exportar_indicadores", args=["xlsx"]), filtros)
        hoja = load_workbook(BytesIO(excel.content), read_only=True)["Candidatos"]
        self.assertEqual(hoja.max_row, 3)
        self.assertContains(respuesta, "Periodo: Semanal")

    def test_tendencia_anual_muestra_doce_meses_y_mensual_todos_los_dias(self):
        self.candidato(1, fecha=date(2024, 2, 29))
        anual = self.client.get(reverse("procesos:indicadores"), {"periodo": "anual", "anio": "2024"})
        self.assertEqual(len(anual.context["indicadores"]["graficas"]["tendencia"]), 12)
        mensual = self.client.get(reverse("procesos:indicadores"), {"periodo": "mensual", "mes": "2024-02"})
        self.assertEqual(len(mensual.context["indicadores"]["graficas"]["tendencia"]), 29)
