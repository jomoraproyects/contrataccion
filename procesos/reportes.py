from collections import Counter
from datetime import timedelta
from io import BytesIO
from statistics import mean, median
from xml.sax.saxutils import escape

from django.utils import timezone
from django.db.models import Q

from .calendario import SEGUNDOS_JORNADA_LABORAL
from .models import PLAZOS_ETAPA_DIAS, ProcesoSeleccion, SeguimientoEtapa


def filtrar_procesos_por_fecha(queryset, formulario, incluir_fechas=True):
    if not formulario.is_valid():
        return queryset.none()
    desde = formulario.cleaned_data.get("fecha_desde")
    hasta = formulario.cleaned_data.get("fecha_hasta")
    q = formulario.cleaned_data.get("q", "").strip()
    if q:
        for termino in q.split():
            queryset = queryset.filter(
                Q(nombre__icontains=termino)
                | Q(apellidos__icontains=termino)
                | Q(cedula__icontains=termino)
            )
    if desde and incluir_fechas:
        queryset = queryset.filter(fecha_llegada__gte=desde)
    if hasta and incluir_fechas:
        queryset = queryset.filter(fecha_llegada__lte=hasta)
    estado = formulario.cleaned_data.get("estado")
    etapa = formulario.cleaned_data.get("etapa")
    vacante = formulario.cleaned_data.get("vacante", "").strip()
    if estado:
        queryset = queryset.filter(estado=estado)
    if etapa:
        queryset = queryset.filter(etapa_actual=etapa)
    if vacante:
        queryset = queryset.filter(vacante__icontains=vacante)
    return queryset


def calcular_indicadores(procesos, periodo=None):
    procesos = list(procesos)
    seguimientos = [item for proceso in procesos for item in proceso.seguimientos.all()]
    cerrados = [
        item for item in seguimientos
        if item.fin is not None and item.resultado != SeguimientoEtapa.Resultado.AJUSTADO
    ]
    abiertos = [
        item for item in seguimientos
        if item.fin is None and item.proceso.activo and not item.proceso.esta_finalizado
    ]
    duraciones_cerradas = {item.pk: item.duracion_horas for item in cerrados}

    contratados = sum(p.estado == ProcesoSeleccion.Estado.CONTRATADO for p in procesos)
    rechazados = sum(p.estado == ProcesoSeleccion.Estado.RECHAZADO for p in procesos)
    finalizados = contratados + rechazados

    por_mes_contador = Counter((p.fecha_llegada.year, p.fecha_llegada.month) for p in procesos)
    meses = [
        {
            "periodo": f"{anio:04d}-{mes:02d}",
            "etiqueta": f"{mes:02d}/{anio}",
            "cantidad": cantidad,
        }
        for (anio, mes), cantidad in sorted(por_mes_contador.items(), reverse=True)[:5]
    ]

    metricas_etapa = []
    for valor, etiqueta in ProcesoSeleccion.Etapa.choices:
        cerrados_etapa = [item for item in cerrados if item.etapa == valor]
        todos_etapa = [item for item in seguimientos if item.etapa == valor and item.resultado != SeguimientoEtapa.Resultado.AJUSTADO]
        duraciones = [duraciones_cerradas[item.pk] for item in cerrados_etapa]
        medidos = [item for item in cerrados_etapa] + [item for item in abiertos if item.etapa == valor]
        retrasos_cerrados = sum(item.esta_vencido for item in cerrados_etapa)
        metricas_etapa.append({
            "valor": valor,
            "etapa": etiqueta,
            "promedio_horas": round(mean(duraciones), 1) if duraciones else None,
            "promedio_observado_horas": round(mean(item.duracion_horas for item in medidos), 1) if medidos else None,
            "medidos": len(medidos),
            "mediana_horas": round(median(duraciones), 1) if duraciones else None,
            "finalizadas": len(cerrados_etapa),
            "abiertas": sum(item.fin is None for item in todos_etapa),
            "retrasos": sum(item.esta_vencido for item in todos_etapa),
            "retrasos_cerrados": retrasos_cerrados,
            "a_tiempo": len(cerrados_etapa) - retrasos_cerrados,
            "cumplimiento_porcentaje": round((len(cerrados_etapa) - retrasos_cerrados) * 100 / len(cerrados_etapa)) if cerrados_etapa else None,
            "plazo_dias": PLAZOS_ETAPA_DIAS[valor],
            "plazo_horas": PLAZOS_ETAPA_DIAS[valor] * SEGUNDOS_JORNADA_LABORAL / 3600,
        })

    con_retrasos = [item for item in metricas_etapa if item["retrasos"]]
    area_mas_retrasos = max(con_retrasos, key=lambda item: item["retrasos"], default=None)
    proximos_vencer = sorted(
        [item for item in abiertos if not item.esta_vencido and item.dias_restantes <= 1],
        key=lambda item: item.fecha_limite,
    )
    conteo_vacantes = Counter()
    nombres_vacantes = {}
    for proceso in procesos:
        nombre = proceso.vacante.strip()
        clave = nombre.casefold()
        conteo_vacantes[clave] += 1
        nombres_vacantes.setdefault(clave, nombre)
    todas_vacantes = [
        {"vacante": nombres_vacantes[clave], "cantidad": cantidad}
        for clave, cantidad in conteo_vacantes.most_common()
    ]
    vacantes = todas_vacantes[:5]
    # Porcentajes enteros para un gráfico SVG sin dependencias de JavaScript.
    distribucion = []
    inicio = 0
    for vacante in vacantes[:4]:
        porcentaje = round(vacante["cantidad"] * 100 / len(procesos)) if procesos else 0
        distribucion.append({**vacante, "porcentaje": porcentaje, "inicio": inicio, "resto": 100 - porcentaje})
        inicio += porcentaje
    if len(todas_vacantes) > 4:
        cantidad = sum(item["cantidad"] for item in todas_vacantes[4:])
        distribucion.append({"vacante": "Otras vacantes", "cantidad": cantidad, "porcentaje": max(0, 100 - inicio), "inicio": inicio, "resto": min(100, inicio)})
    elif distribucion:
        distribucion[-1]["porcentaje"] += 100 - inicio
        distribucion[-1]["resto"] = 100 - distribucion[-1]["porcentaje"]
    promedios_etapa = [item["promedio_horas"] for item in metricas_etapa if item["promedio_horas"] is not None]
    max_promedio_horas = max([1, *promedios_etapa])
    for item in metricas_etapa:
        item["barra_porcentaje"] = round(item["promedio_horas"] * 100 / max_promedio_horas) if item["promedio_horas"] is not None else 0
    area_mas_demorada = max(
        (item for item in metricas_etapa if item["promedio_horas"] is not None),
        key=lambda item: item["promedio_horas"], default=None,
    )
    area_mayor_tiempo_observado = max(
        (item for item in metricas_etapa if item["promedio_observado_horas"] is not None),
        key=lambda item: item["promedio_observado_horas"], default=None,
    )

    # Serie mensual continua: los meses sin postulaciones aparecen como cero.
    # Se muestran hasta doce meses para conservar etiquetas legibles en el gráfico.
    tendencia = []
    if por_mes_contador:
        inicio_mes = min(anio * 12 + mes - 1 for anio, mes in por_mes_contador)
        fin_mes = max(anio * 12 + mes - 1 for anio, mes in por_mes_contador)
        for indice in range(max(inicio_mes, fin_mes - 11), fin_mes + 1):
            anio, mes_cero = divmod(indice, 12)
            mes = mes_cero + 1
            tendencia.append({"etiqueta": f"{mes:02d}/{anio}", "cantidad": por_mes_contador[(anio, mes)]})
    tendencia_titulo = "Procesos iniciados por mes"
    tendencia_descripcion = "Hasta doce meses del periodo filtrado. Los meses sin candidatos aparecen en cero."
    if periodo and periodo.get("periodo") in ("semanal", "mensual", "anual"):
        tipo = periodo["periodo"]
        desde, hasta = periodo["fecha_desde"], periodo["fecha_hasta"]
        tendencia = []
        if tipo == "anual":
            tendencia = [{"etiqueta": f"{mes:02d}/{desde.year}", "cantidad": por_mes_contador[(desde.year, mes)]} for mes in range(1, 13)]
            tendencia_descripcion = "Comparación de los doce meses del año seleccionado."
        else:
            conteo_dias = Counter(p.fecha_llegada for p in procesos)
            cursor = desde
            while cursor <= hasta:
                tendencia.append({"etiqueta": cursor.strftime("%d/%m"), "cantidad": conteo_dias[cursor]})
                cursor += timedelta(days=1)
            tendencia_titulo = "Procesos iniciados por día"
            tendencia_descripcion = "Candidatos recibidos cada día de la semana seleccionada." if tipo == "semanal" else "Candidatos recibidos cada día del mes seleccionado."
    estados = [
        {"codigo": valor, "etiqueta": etiqueta, "cantidad": sum(p.estado == valor for p in procesos)}
        for valor, etiqueta in ProcesoSeleccion.Estado.choices
    ]
    rechazos_etapa = [
        {"etiqueta": etiqueta, "cantidad": sum(p.estado == ProcesoSeleccion.Estado.RECHAZADO and p.etapa_actual == valor for p in procesos)}
        for valor, etiqueta in ProcesoSeleccion.Etapa.choices
    ]
    total_vencidos = sum(item.esta_vencido for item in abiertos)
    cumplimiento = round((len(cerrados) - sum(item.esta_vencido for item in cerrados)) * 100 / len(cerrados)) if cerrados else None

    return {
        "total": len(procesos),
        "en_curso": sum(p.estado == ProcesoSeleccion.Estado.EN_CURSO for p in procesos),
        "listos": sum(p.estado == ProcesoSeleccion.Estado.LISTO_CONTRATACION for p in procesos),
        "contratados": contratados,
        "rechazados": rechazados,
        "finalizados": finalizados,
        "porcentaje_contratados": round(contratados * 100 / finalizados) if finalizados else 0,
        "porcentaje_rechazados": round(rechazados * 100 / finalizados) if finalizados else 0,
        "por_mes": meses,
        "max_por_mes": max((item["cantidad"] for item in meses), default=1),
        "por_etapa": metricas_etapa,
        "max_promedio_horas": max_promedio_horas,
        "area_mas_retrasos": area_mas_retrasos,
        "area_mas_demorada": area_mas_demorada,
        "area_mayor_tiempo_observado": area_mayor_tiempo_observado,
        "proximos_vencer": proximos_vencer,
        "vencidos": total_vencidos,
        "cumplimiento_porcentaje": cumplimiento,
        "etapas_finalizadas": len(cerrados),
        "graficas": {
            "estados": estados,
            "rechazos": rechazos_etapa,
            "vacantes": [{"etiqueta": item["vacante"], "cantidad": item["cantidad"]} for item in vacantes],
            "tendencia": tendencia,
            "tendencia_titulo": tendencia_titulo,
            "tendencia_descripcion": tendencia_descripcion,
            "etapas": metricas_etapa,
            "total": len(procesos),
        },
        "vacantes": vacantes,
        "vacante_mas_candidatos": todas_vacantes[0] if todas_vacantes else None,
        "distribucion_vacantes": distribucion,
        "max_vacantes": max((item["cantidad"] for item in vacantes), default=1),
    }


def _texto_excel(valor):
    texto = str(valor or "")
    return f"'{texto}" if texto.startswith(("=", "+", "-", "@")) else texto


def _datos_rechazo(proceso):
    if proceso.estado != ProcesoSeleccion.Estado.RECHAZADO:
        return "No aplica", "No aplica"
    return proceso.etapa_rechazo, proceso.motivo_rechazo


def _requisitos_rechazo(proceso):
    if proceso.estado != ProcesoSeleccion.Estado.RECHAZADO:
        return "No aplica"
    return ", ".join(proceso.requisitos_rechazo) or "Sin lista de verificación registrada"


def _detalle_rechazos_pdf(procesos, estilos, agrupar=False):
    from reportlab.lib.units import cm
    from reportlab.platypus import KeepTogether, Paragraph, Spacer

    rechazados = [proceso for proceso in procesos if proceso.estado == ProcesoSeleccion.Estado.RECHAZADO]
    if not rechazados:
        return []
    contenido = [Spacer(1, 0.35 * cm), Paragraph("Candidatos rechazados: etapa y motivo", estilos["Heading2"])]
    for proceso in rechazados:
        etapa, motivo = _datos_rechazo(proceso)
        motivo_seguro = escape(motivo).replace("\r\n", "\n").replace("\n", "<br/>")
        bloque = [
            Paragraph(f"<b>{escape(proceso.nombre_completo)}</b> - {escape(proceso.cedula)}", estilos["BodyText"]),
            Paragraph(f"<b>Etapa de rechazo:</b> {escape(etapa)}", estilos["BodyText"]),
            Paragraph(f"<b>Requisitos no cumplidos:</b> {escape(_requisitos_rechazo(proceso))}", estilos["BodyText"]),
            Paragraph(f"<b>Motivo:</b> {motivo_seguro}", estilos["BodyText"]),
            Spacer(1, 0.25 * cm),
        ]
        contenido.extend([KeepTogether(bloque)] if agrupar else bloque)
    return contenido


def exportar_excel(procesos, indicadores, rango):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Candidatos"
    hoja.append([
        "Candidato", "Cédula", "Celular", "Vacante",
        "Fecha de llegada de la hoja de vida", "Etapa actual", "Estado",
        "Etapa de rechazo", "Motivo de rechazo",
        "Requisitos no cumplidos",
    ])
    for proceso in procesos:
        etapa_rechazo, motivo_rechazo = _datos_rechazo(proceso)
        hoja.append([
            _texto_excel(proceso.nombre_completo),
            _texto_excel(proceso.cedula),
            _texto_excel(proceso.celular),
            _texto_excel(proceso.vacante),
            proceso.fecha_llegada,
            proceso.get_etapa_actual_display(),
            proceso.get_estado_display(),
            _texto_excel(etapa_rechazo),
            _texto_excel(motivo_rechazo),
            _texto_excel(_requisitos_rechazo(proceso)),
        ])
    for celda in hoja["E"][1:]:
        celda.number_format = "yyyy-mm-dd"

    encabezado = PatternFill("solid", fgColor="1559A6")
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    hoja.sheet_view.showGridLines = False
    for celda in hoja[1]:
        celda.fill = encabezado
        celda.font = Font(color="FFFFFF", bold=True)
        celda.alignment = Alignment(vertical="center")
    for columna in hoja.columns:
        ancho = min(max(len(str(celda.value or "")) for celda in columna) + 3, 44)
        hoja.column_dimensions[columna[0].column_letter].width = max(ancho, 14)

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def exportar_pdf(procesos, indicadores, rango):
    from .pdf_reportes import candidatos_pdf

    return candidatos_pdf(procesos, rango)


def _fecha_excel(valor):
    if valor is None:
        return None
    if timezone.is_aware(valor):
        return timezone.localtime(valor).replace(tzinfo=None)
    return valor


def exportar_trazabilidad_excel(procesos, resumen, filtros):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    libro = Workbook()
    hoja_resumen = libro.active
    hoja_resumen.title = "Resumen"
    hoja_resumen.append(["Trazabilidad de selección"])
    hoja_resumen.append(["Filtros aplicados", filtros])
    hoja_resumen.append([])
    hoja_resumen.append(["Indicador", "Valor"])
    hoja_resumen.append(["Procesos encontrados", resumen["total"]])
    hoja_resumen.append(["Procesos finalizados", resumen["finalizados"]])
    hoja_resumen.append(["Procesos con retrasos", resumen["con_retrasos"]])
    hoja_resumen.append(["Etapas abiertas", resumen["abiertos"]])

    hoja_candidatos = libro.create_sheet("Candidatos")
    hoja_candidatos.append([
        "ID", "Candidato", "Cédula", "Vacante o área", "Fecha de llegada",
        "Etapa actual", "Estado", "Tiempo total hábil (h)", "Etapa de mayor demora",
        "Etapas fuera de tiempo", "Etapa de rechazo", "Motivo de rechazo", "Requisitos no cumplidos",
    ])
    for proceso in procesos:
        etapa_rechazo, motivo_rechazo = _datos_rechazo(proceso)
        hoja_candidatos.append([
            proceso.pk,
            _texto_excel(proceso.nombre_completo),
            _texto_excel(proceso.cedula),
            _texto_excel(proceso.vacante),
            proceso.fecha_llegada,
            proceso.get_etapa_actual_display(),
            proceso.get_estado_display(),
            proceso.traza_total_horas,
            proceso.traza_mas_lenta.get_etapa_display() if proceso.traza_mas_lenta else "-",
            proceso.traza_fuera_tiempo,
            _texto_excel(etapa_rechazo),
            _texto_excel(motivo_rechazo),
            _texto_excel(_requisitos_rechazo(proceso)),
        ])
    for celda in hoja_candidatos["E"][1:]:
        celda.number_format = "yyyy-mm-dd"

    hoja_etapas = libro.create_sheet("Detalle por etapa")
    hoja_etapas.append([
        "Candidato", "Cédula", "Vacante o área", "Etapa", "Plazo (días hábiles)",
        "Inicio", "Finalización", "Fecha límite", "Tiempo hábil (h)", "Resultado", "Cumplimiento",
    ])
    for proceso in procesos:
        for seguimiento in proceso.traza_seguimientos:
            hoja_etapas.append([
                _texto_excel(proceso.nombre_completo),
                _texto_excel(proceso.cedula),
                _texto_excel(proceso.vacante),
                seguimiento.get_etapa_display(),
                seguimiento.plazo_dias,
                _fecha_excel(seguimiento.inicio),
                _fecha_excel(seguimiento.fin),
                _fecha_excel(seguimiento.fecha_limite),
                seguimiento.duracion_horas,
                seguimiento.get_resultado_display() or "Pendiente",
                seguimiento.estado_plazo,
            ])
    for columna in ("F", "G", "H"):
        for celda in hoja_etapas[columna][1:]:
            celda.number_format = "yyyy-mm-dd hh:mm"


    encabezado = PatternFill("solid", fgColor="1559A6")
    for hoja in libro.worksheets:
        hoja.freeze_panes = "A2"
        hoja.sheet_view.showGridLines = False
        for celda in hoja[1]:
            celda.fill = encabezado
            celda.font = Font(color="FFFFFF", bold=True)
            celda.alignment = Alignment(vertical="center")
        for columna in hoja.columns:
            ancho = min(max(len(str(celda.value or "")) for celda in columna) + 3, 44)
            hoja.column_dimensions[columna[0].column_letter].width = max(ancho, 12)
    hoja_resumen["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    hoja_resumen.freeze_panes = None

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def exportar_trazabilidad_pdf(procesos, resumen, filtros):
    from .pdf_reportes import trazabilidad_pdf

    return trazabilidad_pdf(procesos, resumen, filtros)
