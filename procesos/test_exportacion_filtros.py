from datetime import date, timedelta
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .models import Perfil, ProcesoSeleccion
from .pdf_reportes import _tabla_candidatos
from .requisitos import obtener_requisitos


class ExportacionFiltrosTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.usuarios = {}
        for rol in (Perfil.Rol.VINCULACION, Perfil.Rol.GERENTE, *ProcesoSeleccion.Etapa.values):
            usuario = User.objects.create_user(f"filtros_{rol.lower()}")
            usuario.perfil.rol = rol
            usuario.perfil.save()
            cls.usuarios[rol] = usuario

        def nuevo(nombre, cedula, vacante="Operaciones", fecha=date(2026, 10, 8)):
            return ProcesoSeleccion.objects.create(
                nombre=nombre, apellidos="Reporte", cedula=cedula, celular="3001234567",
                vacante=vacante, fecha_llegada=fecha, creado_por=cls.usuarios[Perfil.Rol.VINCULACION],
            )

        def decidir(proceso, rechazar=False):
            requisitos = {r["codigo"]: "CUMPLE" for r in obtener_requisitos(proceso.etapa_actual)}
            if rechazar:
                requisitos[next(iter(requisitos))] = "NO_CUMPLE"
            proceso.registrar_decision(
                cls.usuarios[proceso.etapa_actual], "RECHAZADO" if rechazar else "APROBADO",
                "Motivo exclusivo del candidato rechazado" if rechazar else "Cumple los requisitos de la etapa",
                verificaciones=requisitos,
            )

        cls.pendientes = [nuevo(f"Pendiente {i}", f"56000{i}") for i in range(7)]
        cls.vencido = cls.pendientes[0]
        decidir(cls.vencido)
        decidir(cls.vencido)
        seguimiento = cls.vencido.seguimientos.get(fin__isnull=True)
        seguimiento.fecha_limite = timezone.now() - timedelta(hours=1)
        seguimiento.save(update_fields=["fecha_limite"])
        cls.rechazado = nuevo("CasoRechazado", "569001")
        decidir(cls.rechazado, rechazar=True)
        cls.listo = nuevo("CasoListo", "569002")
        cls.contratado = nuevo("CasoContratado", "569003")
        for proceso, pasos in ((cls.listo, 3), (cls.contratado, 4)):
            for _ in range(pasos):
                decidir(proceso)
        nuevo("OtraVacante", "569004", vacante="Compras")
        nuevo("OtroPeriodo", "569005", fecha=date(2025, 1, 1))
        cls.incluidos = [*cls.pendientes, cls.rechazado, cls.listo, cls.contratado]

    def setUp(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        self.filtros = {"fecha_desde": "2026-10-01", "fecha_hasta": "2026-10-31", "vacante": "oPeRaCiOnEs"}

    def exportar(self, modulo, formato, filtros):
        if modulo == "trazabilidad":
            return self.client.get(reverse("procesos:trazabilidad"), {**filtros, "exportar": formato})
        return self.client.get(reverse("procesos:exportar_indicadores", args=[formato]), filtros)

    def test_pdf_y_excel_respetan_estados_y_no_limitan_a_la_pagina(self):
        for modulo in ("indicadores", "trazabilidad"):
            for estado, esperados in (("", self.incluidos), ("EN_CURSO", self.pendientes), ("RECHAZADO", [self.rechazado]), ("CONTRATADO", [self.contratado]), ("LISTO_CONTRATACION", [self.listo])):
                with self.subTest(modulo=modulo, estado=estado):
                    filtros = {**self.filtros, "estado": estado, "page": 3}
                    with patch("procesos.pdf_reportes._tabla_candidatos", wraps=_tabla_candidatos) as tabla:
                        pdf = self.exportar(modulo, "pdf", filtros)
                        self.assertEqual(pdf.status_code, 200)
                        self.assertTrue(pdf.content.startswith(b"%PDF"))
                        exportados = tabla.call_args.args[1]
                        self.assertEqual({p.pk for p in exportados}, {p.pk for p in esperados})
                    excel = self.exportar(modulo, "xlsx", filtros)
                    hoja = load_workbook(BytesIO(excel.content), read_only=True)["Candidatos"]
                    cabeceras = [c.value for c in hoja[1]]
                    cedula = cabeceras.index("Cédula")
                    motivo = cabeceras.index("Motivo de rechazo")
                    filas = list(hoja.iter_rows(min_row=2, values_only=True))
                    self.assertEqual({fila[cedula] for fila in filas}, {p.cedula for p in esperados})
                    for fila in filas:
                        self.assertEqual(fila[motivo], "Motivo exclusivo del candidato rechazado" if fila[cedula] == self.rechazado.cedula else "No aplica")

    def test_busqueda_etapa_y_cumplimiento_respetan_los_candidatos(self):
        casos = [({"q": "CasoRechazado"}, {self.rechazado.pk}), ({"etapa": "CONTRATACION"}, {self.listo.pk, self.contratado.pk})]
        for modulo in ("indicadores", "trazabilidad"):
            for filtros, esperados in casos:
                with self.subTest(modulo=modulo, filtros=filtros):
                    with patch("procesos.pdf_reportes._tabla_candidatos", wraps=_tabla_candidatos) as tabla:
                        respuesta = self.exportar(modulo, "pdf", {**self.filtros, **filtros})
                        self.assertEqual(respuesta.status_code, 200)
                        self.assertEqual({p.pk for p in tabla.call_args.args[1]}, esperados)
        with patch("procesos.pdf_reportes._tabla_candidatos", wraps=_tabla_candidatos) as tabla:
            respuesta = self.exportar("trazabilidad", "pdf", {**self.filtros, "situacion": "vencidos"})
            self.assertEqual(respuesta.status_code, 200)
            self.assertEqual([p.pk for p in tabla.call_args.args[1]], [self.vencido.pk])

    def test_filtros_sin_resultados_exportan_pdf_valido(self):
        for modulo in ("indicadores", "trazabilidad"):
            with patch("procesos.pdf_reportes._tabla_candidatos", wraps=_tabla_candidatos) as tabla:
                respuesta = self.exportar(modulo, "pdf", {**self.filtros, "vacante": "Vacante inexistente"})
                self.assertEqual(respuesta.status_code, 200)
                self.assertEqual(tabla.call_args.args[1], [])
