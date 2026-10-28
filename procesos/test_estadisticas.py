from datetime import date
from io import BytesIO
from zipfile import ZipFile

from django.test import SimpleTestCase
from django.urls import reverse
from openpyxl import load_workbook

from .forms import IndicadoresFiltroForm
from .models import ProcesoSeleccion
from .test_indicadores import IndicadoresPanelTests


class PeriodosEstadisticosTests(SimpleTestCase):
    def test_fechas_personalizadas_requieren_ambos_extremos(self):
        for filtros in ({"periodo": "personalizado"}, {"periodo": "personalizado", "fecha_desde": "2026-10-01"}, {"periodo": "personalizado", "fecha_desde": "2026-10-10", "fecha_hasta": "2026-10-01"}):
            self.assertFalse(IndicadoresFiltroForm(filtros).is_valid())
        form = IndicadoresFiltroForm({"periodo": "personalizado", "fecha_desde": "2026-10-01", "fecha_hasta": "2026-10-07"})
        self.assertTrue(form.is_valid(), form.errors)


class EstadisticasPanelTests(IndicadoresPanelTests):
    def test_historico_respeta_filtros_y_conserva_las_otras_graficas(self):
        self.candidato(1, date(2026, 9, 3), "Ventas")
        self.candidato(2, date(2026, 8, 3), "Compras")
        respuesta = self.client.get(reverse("procesos:indicadores"), {"periodo": "mensual", "mes": "2026-09", "vacante": "ventas"})
        historico = respuesta.context["indicadores"]["historico"]
        self.assertEqual(historico["desde"], "2026-09-01")
        self.assertEqual(historico["filas"][0]["recibidos"], 1)
        self.assertContains(respuesta, 'data-chart="rechazos"')
        self.assertContains(respuesta, 'data-chart="historico"')
        self.assertNotContains(respuesta, 'data-chart="comparativa"')
        for clave in ("estados", "vacantes", "etapas", "cumplimiento"):
            self.assertContains(respuesta, f'data-chart="{clave}"')
        self.assertContains(respuesta, "informe=estadistico")

    def test_rechazos_se_cuentan_una_vez_por_etapa_final(self):
        p = self.candidato(1)
        ProcesoSeleccion.objects.filter(pk=p.pk).update(estado="RECHAZADO", etapa_actual="PSICOLOGIA")
        respuesta = self.client.get(reverse("procesos:indicadores"), {"periodo": "mensual", "mes": "2026-09"})
        datos = respuesta.context["indicadores"]["graficas"]
        self.assertEqual(sum(x["cantidad"] for x in datos["rechazos"]), 1)
        self.assertEqual(datos["rechazos"][1]["cantidad"], 1)

    def test_excel_estadistico_contiene_datos_graficas_y_candidatos_filtrados(self):
        self.candidato(1, date(2026, 9, 3), "Ventas")
        self.candidato(2, date(2026, 8, 3), "Ventas")
        self.candidato(3, date(2026, 9, 4), "Compras")
        filtros = {"periodo": "mensual", "mes": "2026-09", "vacante": "ventas", "informe": "estadistico"}
        respuesta = self.client.get(reverse("procesos:exportar_indicadores", args=["xlsx"]), filtros)
        self.assertEqual(respuesta.status_code, 200)
        libro = load_workbook(BytesIO(respuesta.content))
        self.assertEqual(libro.sheetnames, ["Indicadores", "Histórico", "Etapas", "Candidatos"])
        self.assertEqual(libro["Candidatos"].max_row, 2)
        self.assertEqual(libro["Indicadores"]["B5"].value, 1)
        self.assertIn("ventas", libro["Indicadores"]["B2"].value)
        self.assertEqual(len(libro["Indicadores"]._charts), 1)
        self.assertEqual(len(libro["Histórico"]._charts), 1)
        self.assertEqual(libro["Histórico"]["B7"].value, 1)
        with ZipFile(BytesIO(respuesta.content)) as archivo:
            self.assertEqual(len([n for n in archivo.namelist() if n.startswith("xl/charts/chart") and n.endswith(".xml")]), 2)

    def test_pdf_estadistico_y_exportacion_vacia_son_validos(self):
        for mes in ("2026-09", "2026-10"):
            self.candidato(1 if mes.endswith("09") else 2, date(2026, 9, 3))
            respuesta = self.client.get(reverse("procesos:exportar_indicadores", args=["pdf"]), {"periodo": "mensual", "mes": mes, "informe": "estadistico"})
            self.assertEqual(respuesta.status_code, 200)
            self.assertTrue(respuesta.content.startswith(b"%PDF"))

    def test_informe_invalido_no_exporta(self):
        respuesta = self.client.get(reverse("procesos:exportar_indicadores", args=["xlsx"]), {"periodo": "personalizado", "fecha_desde": "2026-10-05", "informe": "estadistico"})
        self.assertEqual(respuesta.status_code, 400)
