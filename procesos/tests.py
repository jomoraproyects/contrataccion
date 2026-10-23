from datetime import date, datetime, timedelta

from django.contrib.auth.models import User
from django.contrib.auth.hashers import identify_hasher, make_password
from django.conf import settings
from django.core.cache import cache
from django.core.management import call_command
from django.core.management.base import CommandError
from django.core.exceptions import ValidationError
from django.db import IntegrityError, connection, transaction
from django.test import TestCase, override_settings
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone

from .calendario import es_dia_habil, segundos_habiles_entre, sumar_dias_habiles, sumar_segundos_habiles
from .models import Candidato, CambioGerencia, Decision, EventoSeguridad, Perfil, ProcesoSeleccion, SeguimientoEtapa
from .requisitos import obtener_requisitos
from .signals import crear_perfil, iniciar_control_de_tiempo


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class FlujoProcesoTests(TestCase):
    def _verificaciones(self, etapa, resultado=Decision.Resultado.APROBADO):
        verificaciones = {
            requisito["codigo"]: "CUMPLE" for requisito in obtener_requisitos(etapa)
        }
        if resultado == Decision.Resultado.RECHAZADO and verificaciones:
            verificaciones[next(iter(verificaciones))] = "NO_CUMPLE"
        return verificaciones

    def _registrar_decision(self, proceso, usuario, resultado, observacion):
        return proceso.registrar_decision(
            usuario, resultado, observacion,
            verificaciones=self._verificaciones(proceso.etapa_actual, resultado),
        )

    def _datos_decision(self, etapa, resultado, observacion):
        return {
            "resultado": resultado,
            "observacion": observacion,
            **{
                f"requisito_{codigo}": respuesta
                for codigo, respuesta in self._verificaciones(etapa, resultado).items()
            },
        }

    def setUp(self):
        cache.clear()
        self.usuarios = {}
        for rol in Perfil.Rol.values:
            usuario = User.objects.create_user(username=rol.lower(), password="prueba-segura")
            usuario.perfil.rol = rol
            usuario.perfil.save()
            self.usuarios[rol] = usuario
        self.proceso = ProcesoSeleccion.objects.create(
            nombre="Ana", apellidos="Pérez", cedula="12345", celular="3001234567",
            fecha_llegada=date.today(), vacante="Contabilidad",
            creado_por=self.usuarios[Perfil.Rol.VINCULACION],
        )

    def test_flujo_completo_hasta_contratado(self):
        for rol, etapa_esperada in [
            (Perfil.Rol.RRHH, ProcesoSeleccion.Etapa.PSICOLOGIA),
            (Perfil.Rol.PSICOLOGIA, ProcesoSeleccion.Etapa.SEGURIDAD),
            (Perfil.Rol.SEGURIDAD, ProcesoSeleccion.Etapa.CONTRATACION),
        ]:
            self._registrar_decision(self.proceso, self.usuarios[rol], Decision.Resultado.APROBADO, "Cumple los requisitos")
            self.assertEqual(self.proceso.etapa_actual, etapa_esperada)
        self.assertEqual(self.proceso.estado, ProcesoSeleccion.Estado.LISTO_CONTRATACION)
        self._registrar_decision(self.proceso, self.usuarios[Perfil.Rol.CONTRATACION], Decision.Resultado.APROBADO, "Contrato firmado")
        self.assertEqual(self.proceso.estado, ProcesoSeleccion.Estado.CONTRATADO)
        self.assertEqual(self.proceso.decisiones.count(), 4)
        seguimientos = list(self.proceso.seguimientos.order_by("inicio"))
        self.assertEqual(len(seguimientos), 4)
        self.assertEqual([item.plazo_dias for item in seguimientos], [2, 3, 6, 2])
        self.assertTrue(all(item.fin for item in seguimientos))

    def test_aprobar_redirige_al_area_anterior_a_su_bandeja(self):
        self.client.force_login(self.usuarios[Perfil.Rol.RRHH])
        response = self.client.post(reverse("procesos:detalle", args=[self.proceso.pk]), self._datos_decision(
            self.proceso.etapa_actual, Decision.Resultado.APROBADO, "Validación legal finalizada",
        ))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("procesos:lista"))
        self.proceso.refresh_from_db()
        self.assertEqual(self.proceso.etapa_actual, ProcesoSeleccion.Etapa.PSICOLOGIA)
        response = self.client.get(response.url)
        self.assertContains(response, "El proceso fue enviado a Psicología")
        self.assertNotContains(response, "Page not found")

    def test_plazos_se_crean_y_avanzan_con_cada_etapa(self):
        seguimiento = self.proceso.seguimientos.get()
        self.assertEqual(seguimiento.etapa, ProcesoSeleccion.Etapa.RRHH)
        self.assertEqual(seguimiento.fecha_limite, sumar_dias_habiles(seguimiento.inicio, 2))
        self._registrar_decision(
            self.proceso, self.usuarios[Perfil.Rol.RRHH], Decision.Resultado.APROBADO, "Validación completada"
        )
        seguimiento.refresh_from_db()
        self.assertIsNotNone(seguimiento.fin)
        psicologia = self.proceso.seguimientos.get(etapa=ProcesoSeleccion.Etapa.PSICOLOGIA)
        self.assertEqual(psicologia.fecha_limite, sumar_dias_habiles(psicologia.inicio, 3))

    def test_plazo_habil_excluye_festivo_y_fin_de_semana(self):
        inicio = timezone.make_aware(datetime(2026, 8, 6, 9, 0))
        limite = sumar_dias_habiles(inicio, 2)
        self.assertFalse(es_dia_habil(date(2026, 8, 7)))  # Batalla de Boyacá
        self.assertFalse(es_dia_habil(date(2026, 8, 8)))  # Sábado
        self.assertEqual(limite, timezone.make_aware(datetime(2026, 8, 11, 9, 0)))

    def test_tiempo_habil_solo_cuenta_de_siete_a_cuatro(self):
        inicio = timezone.make_aware(datetime(2026, 8, 10, 6, 0))
        fin = timezone.make_aware(datetime(2026, 8, 10, 18, 0))
        self.assertEqual(segundos_habiles_entre(inicio, fin), 9 * 60 * 60)

        inicio_tarde = timezone.make_aware(datetime(2026, 8, 10, 15, 0))
        limite = sumar_dias_habiles(inicio_tarde, 1)
        self.assertEqual(limite, timezone.make_aware(datetime(2026, 8, 11, 15, 0)))

    def test_asignacion_fuera_de_horario_empieza_en_siguiente_jornada(self):
        inicio = timezone.make_aware(datetime(2026, 8, 10, 18, 30))
        limite = sumar_dias_habiles(inicio, 1)
        self.assertEqual(limite, timezone.make_aware(datetime(2026, 8, 11, 16, 0)))

    def test_duracion_legible_usa_jornadas_de_nueve_horas(self):
        seguimiento = self.proceso.seguimientos.get()
        seguimiento.inicio = timezone.make_aware(datetime(2026, 8, 10, 7, 0))
        seguimiento.fin = timezone.make_aware(datetime(2026, 8, 11, 9, 0))
        seguimiento.save(update_fields=["inicio", "fin"])
        self.assertEqual(self.proceso.duracion_total_legible, "1 jornada(s) y 2 horas")

    def test_panel_del_rol_notifica_asignaciones_y_dias_restantes(self):
        self.client.force_login(self.usuarios[Perfil.Rol.RRHH])
        response = self.client.get(reverse("procesos:lista"))
        self.assertContains(response, "Procesos asignados a Talento Humano")
        self.assertContains(response, "Quedan 2 día(s)")
        self.assertContains(response, "1 proceso pendiente")

    def test_alerta_del_rol_aparece_en_cualquier_pantalla(self):
        self.client.force_login(self.usuarios[Perfil.Rol.RRHH])
        response = self.client.get(reverse("procesos:detalle", args=[self.proceso.pk]))
        self.assertContains(response, "Tiene 1 proceso pendiente en Talento Humano")
        self.assertContains(response, "El más urgente es <strong>Ana Pérez</strong>", html=False)
        self.assertContains(response, "Quedan 2 día(s)")

    def test_gerente_consulta_trazabilidad_y_vencimientos(self):
        seguimiento = self.proceso.seguimientos.get()
        seguimiento.fecha_limite = timezone.now() - timedelta(hours=1)
        seguimiento.save(update_fields=["fecha_limite"])
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:trazabilidad"))
        self.assertContains(response, "Trazabilidad de procesos")
        self.assertContains(response, "Vencido hace 1 día(s)")
        self.assertContains(response, "Plazo: 6 días")
        self.assertContains(response, "Trazabilidad por candidato")
        self.assertContains(response, 'class="card trace-process-card"')
        self.assertContains(response, "Abrir historial completo")
        self.assertNotContains(response, "Etapas finalizadas recientemente")

    def test_estilos_principales_tienen_version_para_evitar_cache_antigua(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:trazabilidad"))
        self.assertContains(response, "css/app.css?v=20261023.2")

    def test_trazabilidad_filtra_por_fecha_estado_nombre_y_vacante(self):
        otro = ProcesoSeleccion.objects.create(
            nombre="Luis", apellidos="Antiguo", cedula="987654", celular="3007654321",
            fecha_llegada=date(2025, 1, 10), vacante="Ventas",
            creado_por=self.usuarios[Perfil.Rol.VINCULACION],
        )
        self._registrar_decision(
            otro, self.usuarios[Perfil.Rol.RRHH], Decision.Resultado.RECHAZADO, "No continúa"
        )
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:trazabilidad"), {
            "q": "Ana", "fecha_desde": date.today().isoformat(),
            "estado": ProcesoSeleccion.Estado.EN_CURSO, "vacante": "Contabilidad",
        })
        self.assertContains(response, "Ana Pérez")
        self.assertNotContains(response, "Luis Antiguo")
        self.assertContains(response, "Aplicados")
        self.assertContains(response, "1 resultado")

    def test_trazabilidad_valida_el_rango_de_fechas(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:trazabilidad"), {
            "fecha_desde": "2026-08-10", "fecha_hasta": "2026-08-01",
        })
        self.assertContains(response, "La fecha inicial no puede ser posterior")

    def test_trazabilidad_filtra_procesos_con_retrasos(self):
        seguimiento = self.proceso.seguimientos.get()
        seguimiento.fecha_limite = timezone.now() - timedelta(days=1)
        seguimiento.save(update_fields=["fecha_limite"])
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:trazabilidad"), {"situacion": "vencidos"})
        self.assertContains(response, "Ana Pérez")
        self.assertContains(response, "1 etapa fuera de tiempo")

    def test_gerente_recibe_alerta_global_de_vencimientos(self):
        seguimiento = self.proceso.seguimientos.get()
        seguimiento.fecha_limite = timezone.now() - timedelta(hours=2)
        seguimiento.save(update_fields=["fecha_limite"])
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:usuarios"))
        self.assertContains(response, "Alerta de tiempos para Gerencia")
        self.assertContains(response, "1</strong> etapa vencida", html=False)

    def test_solo_gerente_consulta_tablero_de_trazabilidad(self):
        self.client.force_login(self.usuarios[Perfil.Rol.PSICOLOGIA])
        self.assertEqual(self.client.get(reverse("procesos:trazabilidad")).status_code, 403)

    def test_rechazo_detiene_el_proceso(self):
        self._registrar_decision(self.proceso, self.usuarios[Perfil.Rol.RRHH], Decision.Resultado.RECHAZADO, "No cumple validación")
        self.assertEqual(self.proceso.estado, ProcesoSeleccion.Estado.RECHAZADO)
        with self.assertRaises(ValidationError):
            self._registrar_decision(self.proceso, self.usuarios[Perfil.Rol.RRHH], Decision.Resultado.APROBADO, "Cambiar decisión")

    def test_no_permite_saltar_etapa_ni_observacion_vacia(self):
        with self.assertRaises(ValidationError):
            self._registrar_decision(self.proceso, self.usuarios[Perfil.Rol.PSICOLOGIA], Decision.Resultado.APROBADO, "Apto")
        with self.assertRaises(ValidationError):
            self._registrar_decision(self.proceso, self.usuarios[Perfil.Rol.RRHH], Decision.Resultado.APROBADO, "  ")

    def test_base_de_datos_impide_dos_decisiones_vigentes_en_la_misma_etapa(self):
        Decision.objects.create(
            proceso=self.proceso, etapa=ProcesoSeleccion.Etapa.RRHH,
            resultado=Decision.Resultado.APROBADO, observacion="Primera decisión válida",
            decidido_por=self.usuarios[Perfil.Rol.RRHH],
        )
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Decision.objects.create(
                    proceso=self.proceso, etapa=ProcesoSeleccion.Etapa.RRHH,
                    resultado=Decision.Resultado.RECHAZADO, observacion="Duplicado inválido",
                    decidido_por=self.usuarios[Perfil.Rol.RRHH],
                )

    def test_base_de_datos_impide_dos_seguimientos_abiertos(self):
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                SeguimientoEtapa.objects.create(
                    proceso=self.proceso, etapa=ProcesoSeleccion.Etapa.RRHH, ciclo=2,
                    inicio=timezone.now(), fecha_limite=sumar_dias_habiles(timezone.now(), 2),
                )

    def test_gerente_supervisa_pero_no_decide(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:detalle", args=[self.proceso.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "Guardar decisión")
        response = self.client.post(reverse("procesos:detalle", args=[self.proceso.pk]), self._datos_decision(
            self.proceso.etapa_actual, Decision.Resultado.APROBADO, "Apto",
        ))
        self.assertEqual(response.status_code, 403)

    def test_solo_vinculacion_crea_procesos(self):
        for rol in Perfil.Rol.values:
            self.client.force_login(self.usuarios[rol])
            esperado = 200 if rol == Perfil.Rol.VINCULACION else 403
            with self.subTest(rol=rol):
                self.assertEqual(self.client.get(reverse("procesos:crear")).status_code, esperado)

    def test_contratacion_no_crea_ni_consulta_cedulas_por_post_directo(self):
        self.client.force_login(self.usuarios[Perfil.Rol.CONTRATACION])
        respuesta = self.client.post(reverse("procesos:crear"), {
            "nombre": "Otra", "apellidos": "Persona", "cedula": "454545",
            "celular": "3001234567", "fecha_llegada": date.today().isoformat(),
            "vacante": "Sistemas",
        })
        self.assertEqual(respuesta.status_code, 403)
        self.assertFalse(ProcesoSeleccion.objects.filter(cedula="454545").exists())
        self.assertEqual(self.client.get(reverse("procesos:consultar_candidato_cedula"), {
            "cedula": self.proceso.cedula,
        }).status_code, 403)

    def test_flujo_web_desde_vinculacion_hasta_contratacion_separada(self):
        vinculacion = self.usuarios[Perfil.Rol.VINCULACION]
        self.client.force_login(vinculacion)
        respuesta = self.client.post(reverse("procesos:crear"), {
            "nombre": "Elena", "apellidos": "Nueva", "cedula": "454546",
            "celular": "3001234567", "fecha_llegada": date.today().isoformat(),
            "vacante": "Sistemas",
        })
        proceso = ProcesoSeleccion.objects.get(cedula="454546")
        detalle = reverse("procesos:detalle", args=[proceso.pk])
        self.assertRedirects(respuesta, detalle)
        self.assertEqual(proceso.creado_por, vinculacion)
        self.assertEqual(proceso.etapa_actual, ProcesoSeleccion.Etapa.RRHH)
        self.assertEqual(proceso.seguimientos.count(), 1)

        for rol in (Perfil.Rol.RRHH, Perfil.Rol.PSICOLOGIA, Perfil.Rol.SEGURIDAD, Perfil.Rol.CONTRATACION):
            self.client.force_login(vinculacion)
            self.assertEqual(self.client.post(detalle, self._datos_decision(
                rol, Decision.Resultado.APROBADO, "Intento de saltar permisos",
            )).status_code, 403)
            self.client.force_login(self.usuarios[rol])
            self.assertEqual(self.client.get(detalle).status_code, 200)
            respuesta = self.client.post(detalle, self._datos_decision(
                rol, Decision.Resultado.APROBADO, f"Validación completada por {rol}",
            ))
            self.assertEqual(respuesta.status_code, 302)

        proceso.refresh_from_db()
        self.assertEqual(proceso.estado, ProcesoSeleccion.Estado.CONTRATADO)
        self.assertEqual(list(proceso.decisiones.values_list("etapa", flat=True)), list(ProcesoSeleccion.Etapa.values))
        self.assertEqual(proceso.seguimientos.filter(fin__isnull=False).count(), 4)
        self.client.force_login(vinculacion)
        self.assertContains(self.client.get(reverse("procesos:contratacion")), "Elena Nueva")

    def test_contratacion_solo_recibe_el_proceso_tras_aprobar_seguridad(self):
        self.client.force_login(self.usuarios[Perfil.Rol.CONTRATACION])
        detalle = reverse("procesos:detalle", args=[self.proceso.pk])
        self.assertEqual(self.client.get(detalle).status_code, 404)
        self.assertNotContains(self.client.get(reverse("procesos:lista")), "Ana Pérez")
        for rol in (Perfil.Rol.RRHH, Perfil.Rol.PSICOLOGIA, Perfil.Rol.SEGURIDAD):
            self._registrar_decision(self.proceso, self.usuarios[rol], Decision.Resultado.APROBADO, "Validación correcta")
        respuesta = self.client.get(reverse("procesos:lista"))
        self.assertContains(respuesta, "Ana Pérez")
        self.assertContains(respuesta, "Nuevo proceso asignado")
        self.assertNotContains(respuesta, "+ Añadir nuevo proceso")
        self.assertEqual(self.client.get(detalle).status_code, 200)
        self._registrar_decision(
            self.proceso, self.usuarios[Perfil.Rol.CONTRATACION], Decision.Resultado.RECHAZADO, "No acepta condiciones del contrato",
        )
        self.assertContains(self.client.get(reverse("procesos:rechazados")), "No acepta condiciones del contrato")

    def test_vinculacion_ve_sus_procesos_pero_no_los_de_otro_creador(self):
        otro = User.objects.create_user(username="otra.vinculacion", password="prueba-segura")
        otro.perfil.rol = Perfil.Rol.VINCULACION
        otro.perfil.save(update_fields=["rol"])
        self.client.force_login(otro)
        self.assertNotContains(self.client.get(reverse("procesos:lista")), "Ana Pérez")
        self.assertEqual(self.client.get(reverse("procesos:detalle", args=[self.proceso.pk])).status_code, 404)

    def test_migracion_conserva_cuentas_claves_procesos_y_trazabilidad(self):
        from importlib import import_module
        from django.apps import apps
        from types import SimpleNamespace

        usuario_final = self.usuarios[Perfil.Rol.CONTRATACION]
        usuario_final.username = "contratacion.final"
        usuario_final.save(update_fields=["username"])
        creador = self.usuarios[Perfil.Rol.VINCULACION]
        creador.perfil.rol = Perfil.Rol.CONTRATACION
        creador.perfil.save(update_fields=["rol"])
        self._registrar_decision(self.proceso, self.usuarios[Perfil.Rol.RRHH], Decision.Resultado.RECHAZADO, "Motivo histórico")
        tablas = (User, ProcesoSeleccion, Candidato, Decision, SeguimientoEtapa)
        antes = {modelo: list(modelo.objects.order_by("pk").values()) for modelo in tablas}
        migracion = import_module("procesos.migrations.0015_separar_vinculacion_contratacion")
        migracion.conservar_cuentas_creadoras_como_vinculacion(
            apps, SimpleNamespace(connection=connection),
        )
        creador.perfil.refresh_from_db()
        self.assertEqual(creador.perfil.rol, Perfil.Rol.VINCULACION)
        for modelo in tablas:
            self.assertEqual(list(modelo.objects.order_by("pk").values()), antes[modelo])

    def test_migracion_renombra_contratacion_a_vinculacion_sin_cambiar_clave_ni_procesos(self):
        from importlib import import_module
        from django.apps import apps
        from types import SimpleNamespace

        usuario_final = self.usuarios[Perfil.Rol.CONTRATACION]
        usuario_final.username = "contratacion.final"
        usuario_final.save(update_fields=["username"])
        creador = self.usuarios[Perfil.Rol.VINCULACION]
        creador.username = "contratacion"
        creador.save(update_fields=["username"])
        creador.perfil.rol = Perfil.Rol.CONTRATACION
        creador.perfil.save(update_fields=["rol"])
        clave_anterior = creador.password
        pk_anterior = creador.pk
        proceso_anterior = self.proceso.pk

        migracion = import_module("procesos.migrations.0015_separar_vinculacion_contratacion")
        migracion.conservar_cuentas_creadoras_como_vinculacion(
            apps, SimpleNamespace(connection=connection),
        )

        creador.refresh_from_db()
        creador.perfil.refresh_from_db()
        self.assertEqual(creador.username, "vinculacion")
        self.assertEqual(creador.perfil.rol, Perfil.Rol.VINCULACION)
        self.assertEqual(creador.pk, pk_anterior)
        self.assertEqual(creador.password, clave_anterior)
        self.assertTrue(creador.check_password("prueba-segura"))
        self.assertEqual(ProcesoSeleccion.objects.get(pk=proceso_anterior).creado_por_id, pk_anterior)

    def test_migracion_no_pisa_otra_cuenta_vinculacion(self):
        from importlib import import_module
        from django.apps import apps
        from types import SimpleNamespace

        usuario_original = self.usuarios[Perfil.Rol.CONTRATACION]

        migracion = import_module("procesos.migrations.0015_separar_vinculacion_contratacion")
        with self.assertRaisesMessage(RuntimeError, "Ya existe el usuario vinculacion"):
            migracion.conservar_cuentas_creadoras_como_vinculacion(
                apps, SimpleNamespace(connection=connection),
            )
        usuario_original.perfil.refresh_from_db()
        self.assertEqual(usuario_original.username, "contratacion")
        self.assertEqual(usuario_original.perfil.rol, Perfil.Rol.CONTRATACION)

    def test_formulario_rechaza_celular_con_letras_y_fecha_futura(self):
        self.client.force_login(self.usuarios[Perfil.Rol.VINCULACION])
        response = self.client.post(reverse("procesos:crear"), {
            "nombre": "Luis", "apellidos": "Prueba", "cedula": "987654",
            "celular": "300ABC123", "fecha_llegada": (date.today() + timedelta(days=1)).isoformat(),
            "vacante": "Ventas",
        })
        self.assertContains(response, "Escriba un celular válido")
        self.assertContains(response, "no puede estar en el futuro")
        self.assertFalse(ProcesoSeleccion.objects.filter(cedula="987654").exists())

    def test_edicion_muestra_fecha_existente_en_formato_del_navegador(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:proceso_editar", args=[self.proceso.pk]))
        self.assertContains(
            response,
            f'name="fecha_llegada" value="{self.proceso.fecha_llegada:%Y-%m-%d}"',
            html=False,
        )

    def test_busqueda_admite_nombre_completo_y_fecha_invalida(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:lista"), {"q": "Ana Pérez", "fecha_desde": "fecha-invalida"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ana")

    def test_resumen_general_se_muestra(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:lista"))
        self.assertContains(response, "En curso")
        self.assertContains(response, "Listos para contratar")

    def test_cabeceras_basicas_de_seguridad(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:lista"))
        self.assertEqual(response.headers["X-Frame-Options"], "DENY")
        self.assertEqual(response.headers["X-Content-Type-Options"], "nosniff")
        self.assertEqual(response.headers["Referrer-Policy"], "same-origin")
        self.assertIn("default-src 'self'", response.headers["Content-Security-Policy"])
        self.assertIn("no-store", response.headers["Cache-Control"])
        self.assertEqual(response.headers["Permissions-Policy"], "camera=(), microphone=(), geolocation=(), payment=(), usb=()")

    def test_login_bloquea_tras_cinco_intentos(self):
        url = reverse("login")
        for restantes in (4, 3, 2, 1):
            intento = self.client.post(url, {"username": "rrhh", "password": "incorrecta"})
            self.assertEqual(intento.status_code, 200)
            self.assertContains(intento, "No fue posible ingresar")
            self.assertContains(intento, f"Le quedan <strong>{restantes}</strong>", html=False)
        response = self.client.post(url, {"username": "rrhh", "password": "incorrecta"})
        self.assertEqual(response.status_code, 429)
        self.assertContains(response, "Acceso bloqueado temporalmente", status_code=429)
        self.assertEqual(response.headers["Retry-After"], "900")
        perfil = self.usuarios[Perfil.Rol.RRHH].perfil
        perfil.refresh_from_db()
        self.assertEqual(perfil.intentos_fallidos, 5)
        self.assertIsNotNone(perfil.bloqueado_hasta)
        self.assertEqual(self.client.post(url, {"username": "rrhh", "password": "prueba-segura"}).status_code, 429)

    def test_accesos_esenciales_no_se_bloquean_ni_muestran_contador(self):
        url = reverse("login")
        for rol in (Perfil.Rol.GERENTE, Perfil.Rol.VINCULACION, Perfil.Rol.CONTRATACION):
            usuario = self.usuarios[rol]
            for _ in range(6):
                response = self.client.post(url, {
                    "username": usuario.username, "password": "incorrecta",
                })
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, "El usuario o la contraseña no son correctos")
                self.assertNotContains(response, "Le quedan")
                self.assertNotContains(response, "Acceso bloqueado temporalmente")
            usuario.perfil.refresh_from_db()
            self.assertEqual(usuario.perfil.intentos_fallidos, 0)
            self.assertIsNone(usuario.perfil.bloqueado_hasta)

    def test_acceso_esencial_ignora_y_limpia_un_bloqueo_anterior(self):
        usuario = self.usuarios[Perfil.Rol.GERENTE]
        usuario.perfil.intentos_fallidos = 5
        usuario.perfil.bloqueado_hasta = timezone.now() + timedelta(minutes=15)
        usuario.perfil.save()
        response = self.client.post(reverse("login"), {
            "username": usuario.username, "password": "prueba-segura",
        })
        self.assertEqual(response.status_code, 302)
        usuario.perfil.refresh_from_db()
        self.assertEqual(usuario.perfil.intentos_fallidos, 0)
        self.assertIsNone(usuario.perfil.bloqueado_hasta)

    def test_usuario_desconocido_no_revela_si_existe(self):
        response = self.client.post(reverse("login"), {"username": "no-existe", "password": "incorrecta"})
        self.assertContains(response, "El usuario o la contraseña no son correctos")
        self.assertNotContains(response, "no se encuentra registrado")
        self.assertNotContains(response, "intentos antes")

    def test_campos_vacios_no_desgastan_intentos(self):
        url = reverse("login")
        self.assertContains(self.client.post(url, {"username": "", "password": ""}), "No fue posible ingresar")
        self.assertContains(self.client.post(url, {"username": "rrhh", "password": ""}), "No fue posible ingresar")
        perfil = self.usuarios[Perfil.Rol.RRHH].perfil
        perfil.refresh_from_db()
        self.assertEqual(perfil.intentos_fallidos, 0)

    def test_cuenta_inhabilitada_no_revela_su_estado(self):
        usuario = self.usuarios[Perfil.Rol.RRHH]
        usuario.is_active = False
        usuario.save(update_fields=["is_active"])
        response = self.client.post(reverse("login"), {"username": "rrhh", "password": "prueba-segura"})
        self.assertContains(response, "El usuario o la contraseña no son correctos")
        self.assertNotContains(response, "bloqueado por el Gerente")

    def test_login_correcto_limpia_intentos_fallidos(self):
        url = reverse("login")
        self.client.post(url, {"username": "rrhh", "password": "incorrecta"})
        self.client.post(url, {"username": "rrhh", "password": "incorrecta"})
        self.assertEqual(self.client.post(url, {"username": "rrhh", "password": "prueba-segura"}).status_code, 302)
        perfil = self.usuarios[Perfil.Rol.RRHH].perfil
        perfil.refresh_from_db()
        self.assertEqual(perfil.intentos_fallidos, 0)
        self.assertIsNone(perfil.bloqueado_hasta)

    def test_login_incluye_controles_de_accesibilidad(self):
        response = self.client.get(reverse("login"))
        self.assertContains(response, "Saltar al contenido principal")
        self.assertContains(response, "data-theme-toggle")
        self.assertContains(response, 'autocomplete="current-password"')
        self.assertContains(response, "Selección de Personal")
        self.assertContains(response, 'img/serviciudad.jpg')
        self.assertContains(response, 'login-field-user')
        self.assertContains(response, 'login-field-password')
        self.assertContains(response, 'data-password-toggle')
        self.assertNotContains(response, "Información clara")

    def test_usuario_autenticado_no_vuelve_a_ver_el_login(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        self.assertRedirects(self.client.get(reverse("login")), reverse("procesos:lista"))

    def test_panel_admin_fue_retirado(self):
        self.assertEqual(self.client.get("/admin/").status_code, 404)

    def test_gerente_puede_ver_y_crear_usuarios_operativos(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:usuarios"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Talento Humano")
        response = self.client.post(reverse("procesos:usuario_crear"), {
            "username": "nuevo.rrhh", "first_name": "Nuevo", "last_name": "Usuario",
            "rol": Perfil.Rol.RRHH, "password1": "ClaveInicial2026!", "password2": "ClaveInicial2026!",
        })
        self.assertRedirects(response, reverse("procesos:usuarios"))
        nuevo = User.objects.get(username="nuevo.rrhh")
        self.assertEqual(nuevo.perfil.rol, Perfil.Rol.RRHH)
        self.assertTrue(nuevo.check_password("ClaveInicial2026!"))

    def test_gerente_crea_cuenta_final_sin_modificar_a_vinculacion(self):
        creador = self.usuarios[Perfil.Rol.VINCULACION]
        clave_creador = creador.password
        self.usuarios[Perfil.Rol.CONTRATACION].perfil.rol = Perfil.Rol.SIN_ASIGNAR
        self.usuarios[Perfil.Rol.CONTRATACION].perfil.save(update_fields=["rol"])
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        self.assertContains(self.client.get(reverse("procesos:usuarios")), "Falta asignar la última etapa")
        respuesta = self.client.post(reverse("procesos:usuario_crear"), {
            "username": "contratacion.final", "first_name": "Responsable", "last_name": "Final",
            "rol": Perfil.Rol.CONTRATACION,
            "password1": "ClaveFinalElegida2026!", "password2": "ClaveFinalElegida2026!",
        })
        self.assertRedirects(respuesta, reverse("procesos:usuarios"))
        final = User.objects.get(username="contratacion.final")
        self.assertEqual(final.perfil.rol, Perfil.Rol.CONTRATACION)
        self.assertTrue(final.check_password("ClaveFinalElegida2026!"))
        creador.refresh_from_db()
        self.assertEqual(creador.password, clave_creador)
        self.assertEqual(creador.perfil.rol, Perfil.Rol.VINCULACION)
        self.assertNotContains(self.client.get(reverse("procesos:usuarios")), "Falta asignar la última etapa")

    @override_settings(DEBUG=True)
    def test_demo_conserva_contrasenas_y_roles_existentes(self):
        from io import StringIO

        antiguo = self.usuarios[Perfil.Rol.CONTRATACION]
        antiguo.perfil.rol = Perfil.Rol.VINCULACION
        antiguo.perfil.save(update_fields=["rol"])
        clave = antiguo.password
        salida = StringIO()
        call_command("crear_usuarios_demo", password="ClaveDemoNueva2026!", stdout=salida)
        antiguo.refresh_from_db()
        self.assertEqual(antiguo.password, clave)
        self.assertEqual(antiguo.perfil.rol, Perfil.Rol.VINCULACION)
        self.assertIn("se conservan su rol y contraseña", salida.getvalue())

    def test_gerente_edita_y_bloquea_usuario(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        objetivo = self.usuarios[Perfil.Rol.RRHH]
        response = self.client.post(reverse("procesos:usuario_editar", args=[objetivo.pk]), {
            "username": "rrhh.nuevo", "first_name": "María", "last_name": "Recursos Humanos",
            "rol": Perfil.Rol.PSICOLOGIA, "is_active": "",
        })
        self.assertRedirects(response, reverse("procesos:usuarios"))
        objetivo.refresh_from_db()
        objetivo.perfil.refresh_from_db()
        self.assertEqual(objetivo.username, "rrhh.nuevo")
        self.assertEqual(objetivo.get_full_name(), "María Recursos Humanos")
        self.assertEqual(objetivo.perfil.rol, Perfil.Rol.PSICOLOGIA)
        self.assertFalse(objetivo.is_active)
        self.client.logout()
        self.assertFalse(self.client.login(username="rrhh.nuevo", password="prueba-segura"))

    def test_contratacion_no_puede_bloquearse_manualmente(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        objetivo = self.usuarios[Perfil.Rol.CONTRATACION]
        url = reverse("procesos:usuario_editar", args=[objetivo.pk])
        response = self.client.get(url)
        self.assertContains(response, "no puede bloquearse manualmente")
        self.assertTrue(response.context["form"].fields["is_active"].disabled)
        response = self.client.post(url, {
            "username": objetivo.username, "first_name": "Equipo", "last_name": "Contratación",
            "rol": Perfil.Rol.CONTRATACION, "is_active": "",
        })
        self.assertRedirects(response, reverse("procesos:usuarios"))
        objetivo.refresh_from_db()
        self.assertTrue(objetivo.is_active)

    def test_cambiar_un_usuario_a_contratacion_lo_mantiene_activo(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        objetivo = self.usuarios[Perfil.Rol.RRHH]
        response = self.client.post(reverse("procesos:usuario_editar", args=[objetivo.pk]), {
            "username": objetivo.username, "first_name": "Nuevo", "last_name": "Contratación",
            "rol": Perfil.Rol.CONTRATACION, "is_active": "",
        })
        self.assertRedirects(response, reverse("procesos:usuarios"))
        objetivo.refresh_from_db()
        objetivo.perfil.refresh_from_db()
        self.assertTrue(objetivo.is_active)
        self.assertEqual(objetivo.perfil.rol, Perfil.Rol.CONTRATACION)

    def test_modelo_impide_desactivar_accesos_esenciales(self):
        for rol in (Perfil.Rol.GERENTE, Perfil.Rol.VINCULACION, Perfil.Rol.CONTRATACION):
            usuario = self.usuarios[rol]
            usuario.is_active = False
            usuario.save(update_fields=["is_active"])
            usuario.refresh_from_db()
            self.assertTrue(usuario.is_active)

    def test_equipo_oculta_cuentas_sin_rol_operativo(self):
        cuenta_archivada = self.usuarios[Perfil.Rol.SIN_ASIGNAR]
        cuenta_archivada.username = "admin-antiguo"
        cuenta_archivada.save(update_fields=["username"])
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:usuarios"))
        self.assertNotContains(response, "admin-antiguo")

    def test_informacion_del_candidato_usa_filas_completas(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:detalle", args=[self.proceso.pk]))
        self.assertContains(response, 'class="candidate-info-table fs-5"')
        self.assertContains(response, '<th scope="row">', count=6)
        self.assertNotContains(response, 'class="row fs-5 info-list"')

    def test_gerente_cambia_password_operativo(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        objetivo = self.usuarios[Perfil.Rol.PSICOLOGIA]
        objetivo.perfil.intentos_fallidos = 5
        objetivo.perfil.bloqueado_hasta = timezone.now() + timedelta(minutes=15)
        objetivo.perfil.save()
        response = self.client.post(reverse("procesos:usuario_password", args=[objetivo.pk]), {
            "new_password1": "NuevaClave2026!", "new_password2": "NuevaClave2026!",
        })
        self.assertRedirects(response, reverse("procesos:usuarios"))
        objetivo.refresh_from_db()
        objetivo.perfil.refresh_from_db()
        self.assertTrue(objetivo.check_password("NuevaClave2026!"))
        self.assertEqual(objetivo.perfil.intentos_fallidos, 0)
        self.assertIsNone(objetivo.perfil.bloqueado_hasta)

    def test_listado_paginado_sin_perder_filtros(self):
        for indice in range(25):
            ProcesoSeleccion.objects.create(
                nombre=f"Persona {indice}", apellidos="Prueba", cedula=f"99{indice:03d}",
                celular="3001234567", fecha_llegada=date.today(), vacante="Ventas",
                creado_por=self.usuarios[Perfil.Rol.VINCULACION],
            )
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:lista"), {"vacante": "Ventas"})
        self.assertEqual(len(response.context["procesos"]), 5)
        self.assertContains(response, "vacante=Ventas&amp;page=2")
        response = self.client.get(reverse("procesos:lista"), {"vacante": "Ventas", "page": 2})
        self.assertEqual(len(response.context["procesos"]), 5)
        response = self.client.get(reverse("procesos:lista"), {"vacante": "Ventas", "page": 3})
        self.assertEqual(len(response.context["procesos"]), 5)
        self.assertEqual(response.context["page_obj"].paginator.num_pages, 5)

    def test_trazabilidad_muestra_cinco_recientes_por_pagina(self):
        for indice in range(11):
            ProcesoSeleccion.objects.create(
                nombre=f"Reciente {indice:02d}", apellidos="Prueba", cedula=f"881{indice:03d}",
                celular="3001234567", fecha_llegada=date.today(), vacante="Ventas",
                creado_por=self.usuarios[Perfil.Rol.VINCULACION],
            )
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        primera = self.client.get(reverse("procesos:trazabilidad"))
        self.assertEqual(len(primera.context["procesos"]), 5)
        self.assertEqual(primera.context["procesos"][0].nombre, "Reciente 10")
        tercera = self.client.get(reverse("procesos:trazabilidad"), {"page": 3})
        self.assertEqual(len(tercera.context["procesos"]), 2)
        self.assertEqual(tercera.context["page_obj"].paginator.count, 12)

    def test_equipo_muestra_cinco_usuarios_por_pagina(self):
        for indice in range(2):
            usuario = User.objects.create_user(username=f"rrhh_nuevo_{indice}", password="clave-de-prueba")
            usuario.perfil.rol = Perfil.Rol.RRHH
            usuario.perfil.save()
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        primera = self.client.get(reverse("procesos:usuarios"))
        self.assertEqual(len(primera.context["usuarios"]), 5)
        self.assertEqual(primera.context["usuarios"][0].username, "rrhh_nuevo_1")
        segunda = self.client.get(reverse("procesos:usuarios"), {"page": 2})
        self.assertEqual(len(segunda.context["usuarios"]), 3)

    def test_bandeja_area_pagina_pendientes_sin_perder_total(self):
        for indice in range(6):
            ProcesoSeleccion.objects.create(
                nombre=f"Pendiente {indice}", apellidos="Prueba", cedula=f"771{indice:03d}",
                celular="3001234567", fecha_llegada=date.today(), vacante="Ventas",
                creado_por=self.usuarios[Perfil.Rol.VINCULACION],
            )
        self.client.force_login(self.usuarios[Perfil.Rol.RRHH])
        primera = self.client.get(reverse("procesos:lista"))
        self.assertEqual(len(primera.context["asignaciones"]), 5)
        self.assertEqual(primera.context["asignaciones"].paginator.count, 7)
        self.assertContains(primera, "asignaciones_page=2")
        segunda = self.client.get(reverse("procesos:lista"), {"asignaciones_page": 2})
        self.assertEqual(len(segunda.context["asignaciones"]), 2)

    def test_bandeja_area_pagina_historial_terminado(self):
        for indice in range(7):
            proceso = ProcesoSeleccion.objects.create(
                nombre=f"Terminado {indice}", apellidos="Prueba", cedula=f"773{indice:03d}",
                celular="3001234567", fecha_llegada=date.today(), vacante="Ventas",
                creado_por=self.usuarios[Perfil.Rol.VINCULACION],
            )
            self._registrar_decision(
                proceso, self.usuarios[Perfil.Rol.RRHH], Decision.Resultado.RECHAZADO, "No continúa",
            )
        self.client.force_login(self.usuarios[Perfil.Rol.RRHH])
        primera = self.client.get(reverse("procesos:lista"))
        self.assertEqual(len(primera.context["historial_terminados"]), 5)
        self.assertContains(primera, "historial_page=2")
        segunda = self.client.get(reverse("procesos:lista"), {"historial_page": 2})
        self.assertEqual(len(segunda.context["historial_terminados"]), 2)

    def test_usuario_operativo_no_puede_gestionar_cuentas(self):
        self.client.force_login(self.usuarios[Perfil.Rol.RRHH])
        self.assertEqual(self.client.get(reverse("procesos:usuarios")).status_code, 403)
        self.assertEqual(self.client.get(reverse("procesos:usuario_crear")).status_code, 403)

    def test_gerente_no_puede_modificar_otro_gerente(self):
        gerente = self.usuarios[Perfil.Rol.GERENTE]
        otro = User.objects.create_user(username="otro.gerente", password="ClaveGerente2026!")
        otro.perfil.rol = Perfil.Rol.GERENTE
        otro.perfil.save()
        self.client.force_login(gerente)
        self.assertEqual(self.client.get(reverse("procesos:usuario_editar", args=[otro.pk])).status_code, 403)
        self.assertEqual(self.client.get(reverse("procesos:usuario_password", args=[otro.pk])).status_code, 403)

    def test_gerente_edita_datos_sin_poder_saltar_el_flujo(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.post(reverse("procesos:proceso_editar", args=[self.proceso.pk]), {
            "nombre": "Ana María", "apellidos": "Pérez", "cedula": "12345",
            "celular": "3001234567", "fecha_llegada": date.today().isoformat(),
            "vacante": "Contabilidad", "etapa_actual": ProcesoSeleccion.Etapa.CONTRATACION,
            "estado": ProcesoSeleccion.Estado.LISTO_CONTRATACION,
            "motivo_cambio": "Corrección autorizada por documentación recibida.",
        })
        self.assertRedirects(response, reverse("procesos:detalle", args=[self.proceso.pk]))
        self.proceso.refresh_from_db()
        self.assertEqual(self.proceso.nombre, "Ana María")
        self.assertEqual(self.proceso.etapa_actual, ProcesoSeleccion.Etapa.RRHH)
        self.assertEqual(self.proceso.estado, ProcesoSeleccion.Estado.EN_CURSO)
        cambio = self.proceso.cambios_gerencia.get()
        self.assertEqual(cambio.etapa_anterior, ProcesoSeleccion.Etapa.RRHH)
        detalle = self.client.get(reverse("procesos:detalle", args=[self.proceso.pk]))
        self.assertContains(detalle, "Seguridad y Salud en el Trabajo")
        self.assertContains(detalle, "Cambios realizados por Gerencia")

    def test_gerente_no_reabre_un_rechazo_desde_la_edicion(self):
        self._registrar_decision(
            self.proceso, self.usuarios[Perfil.Rol.RRHH], Decision.Resultado.RECHAZADO, "Documento inicialmente inválido"
        )
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        self.client.post(reverse("procesos:proceso_editar", args=[self.proceso.pk]), {
            "nombre": "Ana", "apellidos": "Pérez", "cedula": "12345",
            "celular": "3001234567", "fecha_llegada": date.today().isoformat(),
            "vacante": "Contabilidad", "etapa_actual": ProcesoSeleccion.Etapa.RRHH,
            "estado": ProcesoSeleccion.Estado.EN_CURSO,
            "motivo_cambio": "Se recibió el documento corregido y se reabre la validación.",
        })
        self.proceso.refresh_from_db()
        anterior = self.proceso.decisiones.get()
        self.assertTrue(anterior.vigente)
        self.assertEqual(self.proceso.estado, ProcesoSeleccion.Estado.RECHAZADO)
        seguimientos_rrhh = self.proceso.seguimientos.filter(etapa=ProcesoSeleccion.Etapa.RRHH)
        self.assertEqual(seguimientos_rrhh.count(), 1)
        self.assertEqual(seguimientos_rrhh.filter(fin__isnull=True).count(), 0)
        self.assertEqual(self.proceso.decisiones.filter(etapa=ProcesoSeleccion.Etapa.RRHH, vigente=True).count(), 1)

    def test_usuario_operativo_no_administra_candidatos(self):
        self.client.force_login(self.usuarios[Perfil.Rol.RRHH])
        self.assertEqual(self.client.get(reverse("procesos:proceso_editar", args=[self.proceso.pk])).status_code, 403)
        self.assertEqual(self.client.post(reverse("procesos:proceso_actividad", args=[self.proceso.pk]), {"motivo": "No autorizado"}).status_code, 403)

    def test_gerente_desactiva_y_restaura_candidato(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        url = reverse("procesos:proceso_actividad", args=[self.proceso.pk])
        self.client.post(url, {"motivo": "Registro duplicado en revisión."})
        self.proceso.refresh_from_db()
        self.assertFalse(self.proceso.activo)
        seguimiento = self.proceso.seguimientos.get(fin__isnull=True)
        limite_antes = seguimiento.fecha_limite
        self.proceso.archivado_en = timezone.now() - timedelta(days=2)
        self.proceso.save(update_fields=["archivado_en"])
        pausa_esperada = segundos_habiles_entre(self.proceso.archivado_en, timezone.now())
        limite_esperado = sumar_segundos_habiles(limite_antes, pausa_esperada)
        self.assertNotContains(self.client.get(reverse("procesos:lista")), "Ana Pérez")
        self.assertContains(self.client.get(reverse("procesos:archivados")), "Ana Pérez")
        self.client.post(url, {"motivo": "Se confirmó que el registro es válido."})
        self.proceso.refresh_from_db()
        seguimiento.refresh_from_db()
        self.assertTrue(self.proceso.activo)
        self.assertAlmostEqual((seguimiento.fecha_limite - limite_esperado).total_seconds(), 0, delta=5)
        self.assertAlmostEqual(seguimiento.segundos_pausados, pausa_esperada, delta=5)
        self.assertEqual(self.proceso.cambios_gerencia.count(), 2)

    def test_contratacion_muestra_tiempo_total_y_cumplimiento(self):
        for rol in (Perfil.Rol.RRHH, Perfil.Rol.PSICOLOGIA, Perfil.Rol.SEGURIDAD, Perfil.Rol.CONTRATACION):
            self._registrar_decision(
                self.proceso, self.usuarios[rol], Decision.Resultado.APROBADO, "Etapa completada dentro del plazo"
            )
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:contratacion"))
        self.assertContains(response, "Tiempo total")
        self.assertContains(response, "4 de 4 a tiempo")
        self.assertContains(response, "horas")

    def test_detalle_muestra_inicio_fin_y_tiempo_de_cada_etapa(self):
        self._registrar_decision(
            self.proceso, self.usuarios[Perfil.Rol.RRHH], Decision.Resultado.APROBADO, "Validación terminada"
        )
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:detalle", args=[self.proceso.pk]))
        self.assertContains(response, "Periodo de gestión")
        self.assertContains(response, "Inicio:")
        self.assertContains(response, "Finalizó:")
        self.assertContains(response, "Tiempo empleado")
        self.assertContains(response, "En curso")

    def test_comando_reconstruye_tiempos_de_procesos_antiguos(self):
        for rol in (Perfil.Rol.RRHH, Perfil.Rol.PSICOLOGIA, Perfil.Rol.SEGURIDAD, Perfil.Rol.CONTRATACION):
            self._registrar_decision(
                self.proceso, self.usuarios[rol], Decision.Resultado.APROBADO, "Decisión histórica"
            )
        self.proceso.seguimientos.all().delete()
        call_command("reconstruir_seguimientos")
        self.assertEqual(self.proceso.seguimientos.count(), 4)
        self.assertEqual(self.proceso.seguimientos.filter(fin__isnull=False).count(), 4)

    def test_gerencia_elimina_definitivamente_solo_despues_de_desactivar(self):
        candidato_id = self.proceso.candidato_id
        seguimiento_id = self.proceso.seguimientos.get().pk
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        url_eliminar = reverse("procesos:proceso_eliminar", args=[self.proceso.pk])
        response = self.client.get(url_eliminar)
        self.assertEqual(response.status_code, 404)
        self.client.post(
            reverse("procesos:proceso_actividad", args=[self.proceso.pk]),
            {"motivo": "Registro solicitado para eliminación definitiva."},
        )
        self.assertTrue(ProcesoSeleccion.objects.filter(pk=self.proceso.pk, activo=False).exists())
        self.assertTrue(CambioGerencia.objects.filter(proceso_id=self.proceso.pk).exists())

        response = self.client.get(url_eliminar)
        self.assertContains(response, "Eliminar candidato definitivamente")
        self.assertContains(response, "Ana Pérez")
        self.assertContains(response, "Talento Humano")
        self.assertContains(response, "Esta acción no se puede deshacer")

        self.client.force_login(self.usuarios[Perfil.Rol.RRHH])
        self.assertEqual(self.client.post(url_eliminar, {"confirmar": "on"}).status_code, 403)

        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.post(url_eliminar, {"confirmar": "on"}, follow=True)
        self.assertRedirects(response, reverse("procesos:archivados"))
        self.assertContains(response, "eliminados definitivamente")
        self.assertFalse(ProcesoSeleccion.objects.filter(pk=self.proceso.pk).exists())
        self.assertFalse(SeguimientoEtapa.objects.filter(pk=seguimiento_id).exists())
        self.assertFalse(Candidato.objects.filter(pk=candidato_id).exists())
        self.assertFalse(CambioGerencia.objects.filter(cedula="12345").exists())

    def test_eliminar_exige_confirmacion_explicita(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        self.client.post(
            reverse("procesos:proceso_actividad", args=[self.proceso.pk]),
            {"motivo": "Preparar eliminación con confirmación."},
        )
        response = self.client.post(reverse("procesos:proceso_eliminar", args=[self.proceso.pk]), {})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Este campo es obligatorio")
        self.assertTrue(ProcesoSeleccion.objects.filter(pk=self.proceso.pk).exists())

    def test_eliminar_candidato_exige_desactivar_y_borra_todos_sus_procesos(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        self.client.post(
            reverse("procesos:proceso_actividad", args=[self.proceso.pk]),
            {"motivo": "Desactivar primera postulación para depuración."},
        )
        segundo = ProcesoSeleccion.objects.create(
            nombre="Ana", apellidos="Pérez", cedula="12345", celular="3001234567",
            fecha_llegada=date.today(), vacante="Tesorería",
            creado_por=self.usuarios[Perfil.Rol.VINCULACION],
        )
        url_eliminar = reverse("procesos:proceso_eliminar", args=[self.proceso.pk])
        response = self.client.get(url_eliminar)
        self.assertContains(response, "Hay 1 proceso activo")
        self.assertContains(response, "disabled")
        response = self.client.post(url_eliminar, {"confirmar": "on"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(ProcesoSeleccion.objects.filter(candidato_id=self.proceso.candidato_id).count(), 2)

        self.client.post(
            reverse("procesos:proceso_actividad", args=[segundo.pk]),
            {"motivo": "Desactivar segunda postulación para depuración."},
        )
        response = self.client.post(url_eliminar, {"confirmar": "on"})
        self.assertRedirects(response, reverse("procesos:archivados"))
        self.assertFalse(Candidato.objects.filter(cedula="12345").exists())
        self.assertFalse(ProcesoSeleccion.objects.filter(cedula="12345").exists())
        self.assertFalse(CambioGerencia.objects.filter(cedula="12345").exists())

    def test_cada_area_solo_ve_sus_asignaciones(self):
        self.client.force_login(self.usuarios[Perfil.Rol.PSICOLOGIA])
        self.assertNotContains(self.client.get(reverse("procesos:lista")), "Ana Pérez")
        self.assertEqual(self.client.get(reverse("procesos:detalle", args=[self.proceso.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse("procesos:contratacion")).status_code, 403)

    def test_area_consulta_historial_terminado_en_el_que_participo(self):
        self._registrar_decision(
            self.proceso, self.usuarios[Perfil.Rol.RRHH], Decision.Resultado.APROBADO, "Validación correcta"
        )
        self._registrar_decision(
            self.proceso, self.usuarios[Perfil.Rol.PSICOLOGIA], Decision.Resultado.RECHAZADO, "No continúa"
        )
        self.client.force_login(self.usuarios[Perfil.Rol.RRHH])
        response = self.client.get(reverse("procesos:lista"))
        self.assertContains(response, "Procesos terminados de su área")
        self.assertContains(response, "Ana Pérez")
        self.assertEqual(self.client.get(reverse("procesos:detalle", args=[self.proceso.pk])).status_code, 200)

    def test_no_permite_dos_procesos_abiertos_para_la_misma_cedula(self):
        self.client.force_login(self.usuarios[Perfil.Rol.VINCULACION])
        response = self.client.post(reverse("procesos:crear"), {
            "nombre": "Ana", "apellidos": "Nueva", "cedula": "12345",
            "celular": "3009876543", "fecha_llegada": date.today().isoformat(), "vacante": "Ventas",
        })
        self.assertContains(response, "ya tiene un proceso abierto")
        self.assertEqual(ProcesoSeleccion.objects.filter(cedula="12345").count(), 1)

    def test_base_de_datos_protege_la_cedula_abierta_ante_concurrencia(self):
        with self.assertRaises(IntegrityError):
            ProcesoSeleccion.objects.create(
                nombre="Ana", apellidos="Duplicada", cedula="12345", celular="3009876543",
                fecha_llegada=date.today(), vacante="Ventas",
                creado_por=self.usuarios[Perfil.Rol.VINCULACION],
            )

    def test_nuevo_candidato_rechazado_se_bloquea_con_etapa_y_motivo(self):
        self._registrar_decision(
            self.proceso, self.usuarios[Perfil.Rol.RRHH], Decision.Resultado.RECHAZADO, "Antecedentes sin validar"
        )
        self.client.force_login(self.usuarios[Perfil.Rol.VINCULACION])
        response = self.client.post(reverse("procesos:crear"), {
            "nombre": "Ana", "apellidos": "Pérez", "cedula": "12345",
            "celular": "3009876543", "fecha_llegada": date.today().isoformat(), "vacante": "Ventas",
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "fue rechazada en Talento Humano")
        self.assertContains(response, "Antecedentes sin validar")
        self.assertEqual(ProcesoSeleccion.objects.filter(cedula="12345").count(), 1)
        self.assertEqual(Candidato.objects.filter(cedula="12345").count(), 1)

    def test_nuevo_candidato_con_proceso_contratado_tambien_se_bloquea(self):
        for rol in (
            Perfil.Rol.RRHH, Perfil.Rol.PSICOLOGIA,
            Perfil.Rol.SEGURIDAD, Perfil.Rol.CONTRATACION,
        ):
            self._registrar_decision(self.proceso, self.usuarios[rol], Decision.Resultado.APROBADO, "Apto")
        self.client.force_login(self.usuarios[Perfil.Rol.VINCULACION])
        response = self.client.post(reverse("procesos:crear"), {
            "nombre": "Ana", "apellidos": "Pérez", "cedula": "12345",
            "celular": "3009876543", "fecha_llegada": date.today().isoformat(), "vacante": "Ventas",
        })
        self.assertContains(response, "ya tiene un proceso registrado")
        self.assertContains(response, "Contratado")
        self.assertEqual(ProcesoSeleccion.objects.filter(cedula="12345").count(), 1)

    def test_bloqueo_de_nuevos_candidatos_no_impide_edicion_de_gerencia(self):
        self._registrar_decision(
            self.proceso, self.usuarios[Perfil.Rol.RRHH], Decision.Resultado.RECHAZADO, "No apta",
        )
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.post(reverse("procesos:proceso_editar", args=[self.proceso.pk]), {
            "nombre": "Ana María", "apellidos": "Pérez", "cedula": "12345",
            "celular": "3001234567", "fecha_llegada": date.today().isoformat(),
            "vacante": "Contabilidad", "motivo_cambio": "Corrección del nombre registrado",
        })
        self.assertRedirects(response, reverse("procesos:detalle", args=[self.proceso.pk]))
        self.proceso.refresh_from_db()
        self.assertEqual(self.proceso.nombre, "Ana María")

    def test_vacante_nueva_se_guarda_normalizada(self):
        self.client.force_login(self.usuarios[Perfil.Rol.VINCULACION])
        response = self.client.post(reverse("procesos:crear"), {
            "nombre": "Carlos", "apellidos": "López", "cedula": "778899",
            "celular": "3009876543", "fecha_llegada": date.today().isoformat(),
            "vacante": "  aseo   -   recolección  ",
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(ProcesoSeleccion.objects.get(cedula="778899").vacante, "ASEO - RECOLECCIÓN")

    def test_consulta_cedula_recupera_datos_e_historial(self):
        self.client.force_login(self.usuarios[Perfil.Rol.VINCULACION])
        response = self.client.get(reverse("procesos:consultar_candidato_cedula"), {"cedula": "12.345"})
        self.assertEqual(response.status_code, 200)
        datos = response.json()
        self.assertTrue(datos["encontrado"])
        self.assertTrue(datos["proceso_abierto"])
        self.assertEqual(datos["nombre"], "Ana")
        self.assertEqual(datos["cantidad_procesos"], 1)
        self.assertEqual(datos["ultimo_proceso"]["vacante"], "Contabilidad")
        self.assertIsNone(datos["ultimo_rechazo"])

    def test_consulta_cedula_muestra_el_rechazo_aunque_exista_otro_proceso(self):
        self._registrar_decision(
            self.proceso, self.usuarios[Perfil.Rol.RRHH], Decision.Resultado.RECHAZADO,
            "No aprobó la validación legal",
        )
        ProcesoSeleccion.objects.create(
            nombre="Ana", apellidos="Pérez", cedula="12345", celular="3001234567",
            fecha_llegada=date.today(), vacante="Ventas",
            creado_por=self.usuarios[Perfil.Rol.VINCULACION],
        )
        self.client.force_login(self.usuarios[Perfil.Rol.VINCULACION])
        response = self.client.get(reverse("procesos:consultar_candidato_cedula"), {"cedula": "12.345"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("no-cache", response["Cache-Control"])
        datos = response.json()
        self.assertTrue(datos["proceso_abierto"])
        self.assertEqual(datos["ultimo_rechazo"], {
            "vacante": "Contabilidad", "etapa": "Talento Humano",
            "motivo": "No aprobó la validación legal",
            "requisitos": ["Antecedentes judiciales"],
        })

    def test_consulta_cedula_solo_es_para_vinculacion(self):
        self.client.force_login(self.usuarios[Perfil.Rol.RRHH])
        response = self.client.get(reverse("procesos:consultar_candidato_cedula"), {"cedula": "12345"})
        self.assertEqual(response.status_code, 403)

    def test_formulario_nuevo_incluye_alerta_de_cedula_existente(self):
        self.client.force_login(self.usuarios[Perfil.Rol.VINCULACION])
        response = self.client.get(reverse("procesos:crear"))
        self.assertContains(response, "data-candidate-lookup-url")
        self.assertContains(response, "data-candidate-match")
        self.assertContains(response, "data-candidate-rejection-reason")

    def test_consulta_cedula_rechaza_caracteres_no_permitidos_y_post(self):
        self.client.force_login(self.usuarios[Perfil.Rol.VINCULACION])
        url = reverse("procesos:consultar_candidato_cedula")
        self.assertFalse(self.client.get(url, {"cedula": "abc12345"}).json()["encontrado"])
        self.assertEqual(self.client.post(url, {"cedula": "12345"}).status_code, 405)

    def test_gerencia_no_puede_combinar_historiales_cambiando_la_cedula(self):
        ProcesoSeleccion.objects.create(
            nombre="Carlos", apellidos="Rojas", cedula="987654", celular="3101234567",
            fecha_llegada=date.today(), vacante="Sistemas",
            creado_por=self.usuarios[Perfil.Rol.VINCULACION],
        )
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.post(reverse("procesos:proceso_editar", args=[self.proceso.pk]), {
            "nombre": "Ana", "apellidos": "Pérez", "cedula": "987654",
            "celular": "3001234567", "fecha_llegada": date.today().isoformat(),
            "vacante": "Contabilidad", "motivo_cambio": "Corrección solicitada",
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "no puede combinarse con este historial")
        self.proceso.refresh_from_db()
        self.assertEqual(self.proceso.cedula, "12345")

    def test_observacion_tiene_limite_para_evitar_cargas_excesivas(self):
        self.client.force_login(self.usuarios[Perfil.Rol.RRHH])
        response = self.client.post(reverse("procesos:detalle", args=[self.proceso.pk]), self._datos_decision(
            self.proceso.etapa_actual, Decision.Resultado.APROBADO, "x" * 2001,
        ))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Asegúrese de que este valor tenga como máximo 2000 caracteres")
        self.assertFalse(self.proceso.decisiones.exists())

    def test_identidad_de_candidato_impide_cedulas_duplicadas(self):
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Candidato.objects.create(
                    nombre="Otra", apellidos="Persona", cedula="12345", celular="3009999999"
                )

    def test_carga_de_datos_no_duplica_perfiles_ni_seguimientos(self):
        perfiles_antes = Perfil.objects.count()
        seguimientos_antes = self.proceso.seguimientos.count()
        crear_perfil(sender=User, instance=self.usuarios[Perfil.Rol.RRHH], created=True, raw=True)
        iniciar_control_de_tiempo(
            sender=ProcesoSeleccion, instance=self.proceso, created=True, raw=True
        )
        self.assertEqual(Perfil.objects.count(), perfiles_antes)
        self.assertEqual(self.proceso.seguimientos.count(), seguimientos_antes)

    def test_busqueda_trata_intento_sql_como_texto(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:lista"), {"q": "' OR 1=1 --"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "No hay candidatos")

    def test_datos_del_candidato_se_escapan_contra_xss(self):
        self.proceso.nombre = "<script>alert(1)</script>"
        self.proceso.save(update_fields=["nombre"])
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:detalle", args=[self.proceso.pk]))
        self.assertNotContains(response, "<script>alert(1)</script>")
        self.assertContains(response, "&lt;script&gt;alert(1)&lt;/script&gt;")

    def test_salud_verifica_base_sin_autenticacion(self):
        response = self.client.get(reverse("salud"))
        self.assertEqual(response.status_code, 200)
        self.assertJSONEqual(response.content, {"estado": "ok"})

    def test_login_y_cambios_de_cuenta_quedan_auditados(self):
        self.client.post(reverse("login"), {"username": "rrhh", "password": "incorrecta"})
        self.assertTrue(EventoSeguridad.objects.filter(tipo=EventoSeguridad.Tipo.LOGIN_FALLIDO).exists())
        self.client.post(reverse("login"), {"username": "rrhh", "password": "prueba-segura"})
        self.assertTrue(EventoSeguridad.objects.filter(tipo=EventoSeguridad.Tipo.LOGIN_CORRECTO).exists())

    def test_auditoria_se_conserva_internamente_sin_exponer_una_pantalla(self):
        EventoSeguridad.objects.create(tipo=EventoSeguridad.Tipo.LOGIN_FALLIDO, detalle="Prueba")
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get("/auditoria-seguridad/")
        self.assertEqual(response.status_code, 404)
        self.assertTrue(EventoSeguridad.objects.filter(tipo=EventoSeguridad.Tipo.LOGIN_FALLIDO).exists())

    @override_settings(DEBUG=False)
    def test_404_de_produccion_no_expone_rutas_internas(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get("/ruta-que-no-existe/")
        self.assertEqual(response.status_code, 404)
        self.assertContains(response, "No encontramos esta página", status_code=404)
        self.assertNotContains(response, "Using the URLconf", status_code=404)

    def test_pagina_500_de_produccion_es_generica(self):
        contenido = render_to_string("500.html")
        self.assertIn("No pudimos completar la operación", contenido)
        self.assertNotIn("Traceback", contenido)

    def test_limite_por_ip_detiene_ataques_distribuidos_a_cuentas(self):
        url = reverse("login")
        for indice in range(20):
            response = self.client.post(url, {"username": f"desconocido-{indice}", "password": "incorrecta"})
            self.assertEqual(response.status_code, 200)
        response = self.client.post(url, {"username": "otro", "password": "incorrecta"})
        self.assertEqual(response.status_code, 429)

    def test_cuenta_sin_rol_no_puede_iniciar_sesion(self):
        usuario = self.usuarios[Perfil.Rol.SIN_ASIGNAR]
        response = self.client.post(reverse("login"), {"username": usuario.username, "password": "prueba-segura"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "No fue posible ingresar")
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_todas_las_pantallas_de_gerencia_renderizan(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        urls = [
            reverse("procesos:lista"),
            reverse("procesos:contratacion"),
            reverse("procesos:rechazados"),
            reverse("procesos:archivados"),
            reverse("procesos:trazabilidad"),
            reverse("procesos:usuarios"),
            reverse("procesos:detalle", args=[self.proceso.pk]),
            reverse("procesos:proceso_editar", args=[self.proceso.pk]),
            reverse("procesos:proceso_actividad", args=[self.proceso.pk]),
            reverse("procesos:usuario_crear"),
            reverse("procesos:usuario_editar", args=[self.usuarios[Perfil.Rol.RRHH].pk]),
            reverse("procesos:usuario_password", args=[self.usuarios[Perfil.Rol.RRHH].pk]),
        ]
        for url in urls:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)

    def test_pantallas_operativas_renderizan_segun_el_rol(self):
        casos = [
            (Perfil.Rol.VINCULACION, [
                reverse("procesos:lista"), reverse("procesos:contratacion"),
                reverse("procesos:rechazados"), reverse("procesos:crear"),
                reverse("procesos:detalle", args=[self.proceso.pk]),
            ]),
            (Perfil.Rol.CONTRATACION, [
                reverse("procesos:lista"), reverse("procesos:contratacion"),
                reverse("procesos:rechazados"),
            ]),
            (Perfil.Rol.RRHH, [
                reverse("procesos:lista"), reverse("procesos:detalle", args=[self.proceso.pk]),
            ]),
            (Perfil.Rol.PSICOLOGIA, [reverse("procesos:lista")]),
            (Perfil.Rol.SEGURIDAD, [reverse("procesos:lista")]),
        ]
        for rol, urls in casos:
            self.client.force_login(self.usuarios[rol])
            for url in urls:
                with self.subTest(rol=rol, url=url):
                    self.assertEqual(self.client.get(url).status_code, 200)

    def test_centro_notificaciones_muestra_nuevos_y_permite_marcar_leida(self):
        self.client.force_login(self.usuarios[Perfil.Rol.RRHH])
        response = self.client.get(reverse("procesos:lista"))
        seguimiento = self.proceso.seguimientos.get()
        self.assertContains(response, "Nuevo proceso asignado")
        self.assertContains(response, 'class="notification-badge">1</span>', html=False)

        response = self.client.post(
            reverse("procesos:notificacion_leida", args=[seguimiento.pk]),
            {"next": reverse("procesos:lista")},
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "0 sin leer")
        self.assertNotContains(response, 'class="notification-badge">1</span>', html=False)

    def test_centro_notificaciones_avisa_proximo_vencimiento_y_vencido(self):
        seguimiento = self.proceso.seguimientos.get()
        self.client.force_login(self.usuarios[Perfil.Rol.RRHH])
        seguimiento.fecha_limite = timezone.now() + timedelta(hours=2)
        seguimiento.save(update_fields=["fecha_limite"])
        response = self.client.get(reverse("procesos:lista"))
        self.assertContains(response, "Proceso próximo a vencer")

        seguimiento.fecha_limite = timezone.now() - timedelta(hours=2)
        seguimiento.save(update_fields=["fecha_limite"])
        response = self.client.get(reverse("procesos:lista"))
        self.assertContains(response, "Proceso vencido")

    def test_usuario_no_puede_marcar_notificacion_de_otro_rol(self):
        seguimiento = self.proceso.seguimientos.get()
        self.client.force_login(self.usuarios[Perfil.Rol.PSICOLOGIA])
        response = self.client.post(reverse("procesos:notificacion_leida", args=[seguimiento.pk]))
        self.assertEqual(response.status_code, 404)

    def test_indicadores_disponibles_para_gerencia_y_contratacion(self):
        for rol in (Perfil.Rol.GERENTE, Perfil.Rol.VINCULACION, Perfil.Rol.CONTRATACION):
            self.client.force_login(self.usuarios[rol])
            response = self.client.get(reverse("procesos:indicadores"))
            with self.subTest(rol=rol):
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, "Candidatos por vacante")
                self.assertContains(response, "Candidatos filtrados")
                self.assertContains(response, "Estado del proceso")
                self.assertContains(response, "Tiempo promedio por etapa")
                self.assertContains(response, 'data-chart="tendencia"')
                self.assertContains(response, "Vacantes con más candidatos")
                self.assertContains(response, 'name="periodo"')
                self.assertContains(response, "Semanal")
                self.assertContains(response, "Mensual")
                self.assertContains(response, "Anual")
                self.assertContains(response, "Área o etapa actual")
        self.client.force_login(self.usuarios[Perfil.Rol.RRHH])
        self.assertEqual(self.client.get(reverse("procesos:indicadores")).status_code, 403)

    def test_indicadores_y_exportaciones_respetan_el_ambito_del_rol(self):
        otro_creador = User.objects.create_user(username="vinculacion.otra", password="prueba-segura")
        otro_creador.perfil.rol = Perfil.Rol.VINCULACION
        otro_creador.perfil.save(update_fields=["rol"])
        ProcesoSeleccion.objects.create(
            nombre="Otra", apellidos="Persona", cedula="991144", celular="3001234567",
            fecha_llegada=date.today(), vacante="Ventas", creado_por=otro_creador,
        )
        for rol in (Perfil.Rol.RRHH, Perfil.Rol.PSICOLOGIA, Perfil.Rol.SEGURIDAD):
            self._registrar_decision(self.proceso, self.usuarios[rol], Decision.Resultado.APROBADO, "Apto para continuar")
        for rol, total in ((Perfil.Rol.GERENTE, 2), (Perfil.Rol.VINCULACION, 1), (Perfil.Rol.CONTRATACION, 1)):
            self.client.force_login(self.usuarios[rol])
            with self.subTest(rol=rol):
                respuesta = self.client.get(reverse("procesos:indicadores"))
                self.assertEqual(respuesta.status_code, 200)
                self.assertEqual(respuesta.context["indicadores"]["total"], total)
                if rol != Perfil.Rol.GERENTE:
                    self.assertNotContains(respuesta, "Otra Persona")
        for rol in (Perfil.Rol.RRHH, Perfil.Rol.PSICOLOGIA, Perfil.Rol.SEGURIDAD):
            self.client.force_login(self.usuarios[rol])
            with self.subTest(rol=rol):
                self.assertEqual(self.client.get(reverse("procesos:indicadores")).status_code, 403)
                for formato in ("xlsx", "pdf"):
                    self.assertEqual(self.client.get(
                        reverse("procesos:exportar_indicadores", args=[formato]),
                    ).status_code, 403)

    def test_indicadores_filtran_por_fecha_y_calculan_porcentajes(self):
        anterior = ProcesoSeleccion.objects.create(
            nombre="Luis", apellidos="Viejo", cedula="998877", celular="3009998877",
            fecha_llegada=date(2025, 1, 5), vacante="Archivo",
            creado_por=self.usuarios[Perfil.Rol.VINCULACION],
        )
        self._registrar_decision(
            anterior, self.usuarios[Perfil.Rol.RRHH], Decision.Resultado.RECHAZADO, "No continúa"
        )
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:indicadores"), {
            "fecha_desde": date.today().isoformat(), "fecha_hasta": date.today().isoformat(),
        })
        self.assertEqual(response.context["indicadores"]["total"], 1)
        self.assertContains(response, "Ana Pérez")
        self.assertNotContains(response, "Luis Viejo")

        response = self.client.get(reverse("procesos:indicadores"))
        self.assertEqual(response.context["indicadores"]["rechazados"], 0)
        self.assertNotContains(response, "Luis Viejo")

        response = self.client.get(reverse("procesos:indicadores"), {"periodo": "anual", "anio": 2025})
        self.assertEqual(response.context["indicadores"]["total"], 1)
        self.assertEqual(response.context["indicadores"]["rechazados"], 1)
        self.assertEqual(response.context["indicadores"]["porcentaje_rechazados"], 100)
        self.assertContains(response, "Luis Viejo")
        self.assertEqual([proceso.pk for proceso in response.context["procesos"]], [anterior.pk])

    def test_indicadores_filtran_por_estado_etapa_y_vacante(self):
        otro = ProcesoSeleccion.objects.create(
            nombre="Luis", apellidos="Ventas", cedula="998811", celular="3009998811",
            fecha_llegada=date.today(), vacante="Ventas",
            creado_por=self.usuarios[Perfil.Rol.VINCULACION],
        )
        self._registrar_decision(
            otro, self.usuarios[Perfil.Rol.RRHH], Decision.Resultado.RECHAZADO, "No continúa"
        )
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:indicadores"), {
            "estado": ProcesoSeleccion.Estado.EN_CURSO,
            "etapa": ProcesoSeleccion.Etapa.RRHH,
            "vacante": "Contabilidad",
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["indicadores"]["total"], 1)
        self.assertContains(response, "Ana Pérez")
        self.assertNotContains(response, "Luis Ventas")
        self.assertContains(response, "Estado: En curso")
        self.assertContains(response, "Área/etapa: Talento Humano")

    def test_enlaces_anteriores_indicadores_conservan_filtro_vacante_sin_importar_mayusculas(self):
        ProcesoSeleccion.objects.create(
            nombre="Luis", apellidos="Ventas", cedula="998822", celular="3009998822",
            fecha_llegada=date.today(), vacante="Ventas",
            creado_por=self.usuarios[Perfil.Rol.VINCULACION],
        )
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:indicadores"))
        self.assertContains(response, '<input type="text" name="vacante"', html=False)

        filtrado = self.client.get(reverse("procesos:indicadores"), {
            "q": "998822",
            "vacante": "vEnTaS",
        })
        nombres_filtrados = [proceso.nombre_completo for proceso in filtrado.context["procesos"]]
        self.assertEqual(nombres_filtrados, ["Luis Ventas"])
        self.assertContains(filtrado, "Búsqueda: 998822")
        self.assertContains(filtrado, "Filtros aplicados:")

    def test_indicadores_paginan_cinco_y_exportan_todos_los_filtrados(self):
        from io import BytesIO
        from openpyxl import load_workbook

        for indice in range(12):
            ProcesoSeleccion.objects.create(
                nombre=f"Ventas {indice:02d}", apellidos="Prueba", cedula=f"771{indice:03d}",
                celular="3001234567", fecha_llegada=date.today(), vacante="Ventas",
                creado_por=self.usuarios[Perfil.Rol.VINCULACION],
            )
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        parametros = {
            "fecha_desde": date.today().isoformat(),
            "estado": ProcesoSeleccion.Estado.EN_CURSO,
            "etapa": ProcesoSeleccion.Etapa.RRHH,
            "vacante": "Ventas",
        }
        primera = self.client.get(reverse("procesos:indicadores"), parametros)
        self.assertEqual(len(primera.context["procesos"]), 5)
        self.assertEqual(primera.context["page_obj"].paginator.count, 12)
        self.assertContains(primera, "vacante=Ventas&amp;page=2")

        segunda = self.client.get(reverse("procesos:indicadores"), {**parametros, "page": 2})
        self.assertEqual(len(segunda.context["procesos"]), 5)
        tercera = self.client.get(reverse("procesos:indicadores"), {**parametros, "page": 3})
        self.assertEqual(len(tercera.context["procesos"]), 2)

        excel = self.client.get(reverse("procesos:exportar_indicadores", args=["xlsx"]), parametros)
        libro = load_workbook(BytesIO(excel.content), read_only=True, data_only=True)
        self.assertEqual(libro["Candidatos"].max_row, 13)

    def test_indicadores_destacan_vacante_y_etapa_mas_demorada(self):
        for indice, vacante in enumerate(("Ventas", "VENTAS", "Compras")):
            ProcesoSeleccion.objects.create(
                nombre=f"Persona {indice}", apellidos="Prueba", cedula=f"882{indice:03d}",
                celular="3001234567", fecha_llegada=date.today(), vacante=vacante,
                creado_por=self.usuarios[Perfil.Rol.VINCULACION],
            )
        seguimiento = self.proceso.seguimientos.get()
        seguimiento.inicio = timezone.make_aware(datetime(2026, 9, 8, 7, 0))
        seguimiento.fin = timezone.make_aware(datetime(2026, 9, 8, 12, 0))
        seguimiento.resultado = SeguimientoEtapa.Resultado.APROBADO
        seguimiento.save()
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:indicadores"))
        indicadores = response.context["indicadores"]
        self.assertEqual(indicadores["vacante_mas_candidatos"], {"vacante": "VENTAS", "cantidad": 2})
        self.assertEqual(indicadores["area_mas_demorada"]["etapa"], "Talento Humano")
        self.assertEqual(indicadores["area_mas_demorada"]["promedio_horas"], 5.0)
        self.assertContains(response, 'data-chart="vacantes"')
        self.assertContains(response, 'data-chart="etapas"')

    def test_indicadores_calculan_promedio_y_mediana_por_etapa(self):
        seguimiento = self.proceso.seguimientos.get()
        seguimiento.inicio = timezone.make_aware(datetime(2026, 9, 8, 7, 0))
        seguimiento.fin = timezone.make_aware(datetime(2026, 9, 8, 11, 0))
        seguimiento.resultado = SeguimientoEtapa.Resultado.APROBADO
        seguimiento.save(update_fields=["inicio", "fin", "resultado"])
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:indicadores"))
        talento = response.context["indicadores"]["por_etapa"][0]
        self.assertEqual(talento["promedio_horas"], 4.0)
        self.assertEqual(talento["mediana_horas"], 4.0)

    def test_exportaciones_excel_y_pdf_respetan_filtro(self):
        from io import BytesIO
        from openpyxl import load_workbook

        seguimiento = self.proceso.seguimientos.get()
        seguimiento.fecha_limite = timezone.now() + timedelta(hours=2)
        seguimiento.save(update_fields=["fecha_limite"])
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        parametros = {
            "fecha_desde": date.today().isoformat(),
            "fecha_hasta": date.today().isoformat(),
            "estado": ProcesoSeleccion.Estado.EN_CURSO,
            "vacante": "CONTABILIDAD",
        }
        excel = self.client.get(reverse("procesos:exportar_indicadores", args=["xlsx"]), parametros)
        self.assertEqual(excel.status_code, 200)
        self.assertEqual(
            excel["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        libro = load_workbook(BytesIO(excel.content), read_only=True, data_only=True)
        self.assertEqual(libro.sheetnames, ["Candidatos"])
        self.assertEqual(libro["Candidatos"]["A2"].value, "Ana Pérez")
        self.assertEqual(libro["Candidatos"]["E2"].value.date(), date.today())
        self.assertEqual(libro["Candidatos"]["G2"].value, "En curso")
        self.assertIn(
            f"candidatos-en-curso-contabilidad-desde-{date.today().isoformat()}-hasta-{date.today().isoformat()}",
            excel["Content-Disposition"],
        )

        pdf = self.client.get(reverse("procesos:exportar_indicadores", args=["pdf"]), parametros)
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))

    def test_reportes_incluyen_etapa_y_motivo_del_rechazo(self):
        from io import BytesIO

        from openpyxl import load_workbook
        from reportlab.lib.styles import getSampleStyleSheet

        from .reportes import _detalle_rechazos_pdf

        motivo = "No apto en entrevista de validación"
        self._registrar_decision(
            self.proceso, self.usuarios[Perfil.Rol.RRHH], Decision.Resultado.RECHAZADO, motivo,
        )
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        filtros = {"estado": ProcesoSeleccion.Estado.RECHAZADO}

        excel = self.client.get(reverse("procesos:exportar_indicadores", args=["xlsx"]), filtros)
        self.assertEqual(excel.status_code, 200)
        hoja = load_workbook(BytesIO(excel.content), read_only=True, data_only=True)["Candidatos"]
        self.assertEqual(hoja["H1"].value, "Etapa de rechazo")
        self.assertEqual(hoja["I1"].value, "Motivo de rechazo")
        self.assertEqual(hoja["H2"].value, "Talento Humano")
        self.assertEqual(hoja["I2"].value, motivo)

        traza = self.client.get(reverse("procesos:trazabilidad"), {**filtros, "exportar": "xlsx"})
        self.assertEqual(traza.status_code, 200)
        hoja_traza = load_workbook(BytesIO(traza.content), read_only=True, data_only=True)["Candidatos"]
        self.assertEqual(hoja_traza["K2"].value, "Talento Humano")
        self.assertEqual(hoja_traza["L2"].value, motivo)

        bloques = _detalle_rechazos_pdf([self.proceso], getSampleStyleSheet())
        self.assertTrue(any(motivo in getattr(bloque, "text", "") for bloque in bloques))
        for url in (
            reverse("procesos:exportar_indicadores", args=["pdf"]),
            reverse("procesos:trazabilidad"),
        ):
            parametros = {**filtros, "exportar": "pdf"} if url.endswith("/trazabilidad/") else filtros
            pdf = self.client.get(url, parametros)
            self.assertEqual(pdf.status_code, 200)
            self.assertTrue(pdf.content.startswith(b"%PDF"))

    def test_reportes_no_atribuyen_rechazo_a_candidato_en_curso(self):
        from io import BytesIO

        from openpyxl import load_workbook

        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        excel = self.client.get(reverse("procesos:exportar_indicadores", args=["xlsx"]))
        hoja = load_workbook(BytesIO(excel.content), read_only=True, data_only=True)["Candidatos"]
        self.assertEqual(hoja["H2"].value, "No aplica")
        self.assertEqual(hoja["I2"].value, "No aplica")

    def test_trazabilidad_exporta_excel_y_pdf_con_los_filtros(self):
        from io import BytesIO
        from openpyxl import load_workbook

        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        parametros = {
            "fecha_desde": date.today().isoformat(),
            "estado": ProcesoSeleccion.Estado.EN_CURSO,
            "etapa": ProcesoSeleccion.Etapa.RRHH,
            "vacante": "Contab",
        }
        excel = self.client.get(reverse("procesos:trazabilidad"), {**parametros, "exportar": "xlsx"})
        self.assertEqual(excel.status_code, 200)
        self.assertEqual(
            excel["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        libro = load_workbook(BytesIO(excel.content), read_only=True, data_only=True)
        self.assertEqual(libro.sheetnames, ["Resumen", "Candidatos", "Detalle por etapa", "Resumen por área"])
        self.assertEqual(libro["Candidatos"]["B2"].value, "Ana Pérez")
        self.assertIn("Estado: En curso", libro["Resumen"]["B2"].value)

        pdf = self.client.get(reverse("procesos:trazabilidad"), {**parametros, "exportar": "pdf"})
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))

    def test_error_de_exportacion_regresa_al_informe_sin_mostrar_pagina_500(self):
        from unittest.mock import patch

        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        with patch("procesos.views.exportar_excel", side_effect=RuntimeError("dependencia ausente")):
            response = self.client.get(
                reverse("procesos:exportar_indicadores", args=["xlsx"]),
                {"vacante": "Contabilidad"},
            )
        self.assertEqual(response.status_code, 302)
        self.assertIn("/indicadores/?vacante=Contabilidad", response["Location"])

    def test_equipo_muestra_plazos_habiles_por_rol(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:usuarios"))
        self.assertContains(response, "Plazo de gestión")
        self.assertContains(response, "2 días hábiles")
        self.assertContains(response, "3 días hábiles")
        self.assertContains(response, "6 días hábiles")
        self.assertContains(response, "Supervisión")

    def test_exportacion_rechaza_rango_invalido(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:exportar_indicadores", args=["pdf"]), {
            "fecha_desde": "2026-09-10", "fecha_hasta": "2026-09-01",
        })
        self.assertEqual(response.status_code, 400)

    def test_acciones_importantes_incluyen_confirmacion_contextual(self):
        self.client.force_login(self.usuarios[Perfil.Rol.RRHH])
        response = self.client.get(reverse("procesos:detalle", args=[self.proceso.pk]))
        self.assertContains(response, 'data-confirm-mode="reject-only"', html=False)
        self.assertContains(response, 'data-confirm-subject="Ana Pérez"', html=False)
        self.assertContains(response, 'data-confirm-stage="Talento Humano"', html=False)

        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        response = self.client.get(reverse("procesos:proceso_actividad", args=[self.proceso.pk]))
        self.assertContains(response, 'data-confirm-title="Confirmar desactivación"', html=False)
        response = self.client.get(reverse(
            "procesos:usuario_editar", args=[self.usuarios[Perfil.Rol.RRHH].pk]
        ))
        self.assertContains(response, 'data-confirm-subject="rrhh"', html=False)


class ConfiguracionSeguridadTests(TestCase):
    def test_bcrypt_sha256_es_el_hasher_principal(self):
        self.assertEqual(settings.PASSWORD_HASHERS[0], "django.contrib.auth.hashers.BCryptSHA256PasswordHasher")
        self.assertEqual(identify_hasher(make_password("ClaveLarga2026!", hasher="bcrypt_sha256")).algorithm, "bcrypt_sha256")

    @override_settings(DEBUG=False)
    def test_usuarios_demo_no_se_pueden_crear_en_produccion(self):
        with self.assertRaises(CommandError):
            call_command("crear_usuarios_demo", password="ClaveTemporalSoloPruebas2026!")
