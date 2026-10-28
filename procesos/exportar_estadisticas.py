"""Informes agregados: utiliza los mismos resultados que el panel."""
from io import BytesIO


def excel_estadistico(procesos, resumen, filtros):
    # Se conserva el motor de exportación del servidor y el listado ya aprobado.
    from openpyxl import load_workbook
    from openpyxl.chart import BarChart, DoughnutChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from .reportes import exportar_excel, _texto_excel

    libro = load_workbook(BytesIO(exportar_excel(procesos, resumen, filtros)))
    hoja = libro.create_sheet("Indicadores", 0)
    hoja.append(["Selección de Personal - Serviciudad"])
    hoja.append(["Filtros aplicados", _texto_excel(filtros)])
    hoja.append(["Criterio", "Fecha de llegada HDV. Estados actuales; no decisiones tomadas en el periodo."])
    hoja.append(["Indicador", "Valor"])
    for etiqueta, clave in (("Recibidos", "total"), ("Contratados", "contratados"), ("Rechazados", "rechazados"), ("Pendientes", None), ("Etapas abiertas vencidas", "vencidos")):
        hoja.append([etiqueta, resumen[clave] if clave else resumen["en_curso"] + resumen["listos"]])
    hoja.append(["Cumplimiento de etapas finalizadas", resumen["cumplimiento_porcentaje"] / 100 if resumen["cumplimiento_porcentaje"] is not None else None])
    hoja["B10"].number_format = "0%"
    hoja.append([])
    hoja.append(["Estado", "Candidatos"])
    for item in resumen["graficas"]["estados"]:
        hoja.append([item["etiqueta"], item["cantidad"]])
    dona = DoughnutChart()
    dona.title = "Estado de los candidatos"
    dona.holeSize = 65
    dona.add_data(Reference(hoja, min_col=2, min_row=12, max_row=16), titles_from_data=True)
    dona.set_categories(Reference(hoja, min_col=1, min_row=13, max_row=16))
    dona.width, dona.height = 19, 10
    hoja.add_chart(dona, "D5")
    historico = resumen.get("historico", {})
    hist = libro.create_sheet("Histórico", 1)
    hist.append(["Histórico de procesos", _texto_excel(filtros)])
    hist.append(["Agrupación", historico.get("agrupacion", "mensual")])
    hist.append(["Fechas del histórico", f"{historico.get('desde', '-')} a {historico.get('hasta', '-')}"])
    hist.append(["Criterio", "Recibidos por llegada HDV; contratados y rechazados por fecha de cierre válido, incluso si llegaron antes del rango."])
    hist.append(["Avisos", historico.get("error") or f"{historico.get('sin_fecha', 0)} finalizados sin fecha de cierre verificable; no se inventan fechas."])
    hist.append(["Periodo", "Recibidos", "Contratados", "Rechazados"])
    for fila in historico.get("filas", []):
        hist.append([fila["etiqueta"], fila["recibidos"], fila["contratados"], fila["rechazados"]])
    if historico.get("filas"):
        barras = BarChart()
        barras.title = "Histórico de recibidos, contratados y rechazados"
        barras.y_axis.title = "Cantidad"
        barras.grouping = "clustered"
        barras.overlap = 0
        barras.add_data(Reference(hist, min_col=2, max_col=4, min_row=6, max_row=hist.max_row), titles_from_data=True)
        barras.set_categories(Reference(hist, min_col=1, min_row=7, max_row=hist.max_row))
        barras.width, barras.height = 24, 12
        hist.add_chart(barras, "F7")
    hoja.append([])
    hoja.append(["Vacante (cinco principales)", "Candidatos"])
    for item in resumen["vacantes"]:
        hoja.append([_texto_excel(item["vacante"]), item["cantidad"]])
    etapas = libro.create_sheet("Etapas", 2)
    etapas.append(["Etapa", "Plazo (días hábiles)", "Plazo (h)", "Promedio (h)", "Mediana (h)", "Finalizadas", "A tiempo", "Fuera de plazo", "Cumplimiento", "Abiertas", "Rechazados"])
    for i, etapa in enumerate(resumen["por_etapa"]):
        etapas.append([etapa["etapa"], etapa["plazo_dias"], etapa["plazo_horas"], etapa["promedio_horas"], etapa["mediana_horas"], etapa["finalizadas"], etapa["a_tiempo"], etapa["retrasos_cerrados"], etapa["cumplimiento_porcentaje"] / 100 if etapa["cumplimiento_porcentaje"] is not None else None, etapa["abiertas"], resumen["graficas"]["rechazos"][i]["cantidad"]])
        etapas.cell(i + 2, 9).number_format = "0%"
    etapas.append([])
    etapas.append(["Tiempos", "Solo etapas finalizadas. Se excluyen cierres administrativos de ajuste."])
    etapas.append(["Horario", "7:00 a. m. a 4:00 p. m., lunes a viernes, sin festivos de Colombia."])
    for sheet in (hoja, hist, etapas):
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "B5" if sheet == hoja else "B2"
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F0F5FB")
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1559A6")
            cell.font = Font(color="FFFFFF", bold=True)
        sheet.row_dimensions[1].height = 30
        sheet.column_dimensions["A"].width = 40
        for col in "BCDEFGHIJK":
            sheet.column_dimensions[col].width = 18
    hoja.column_dimensions["B"].width = 52
    hoja.row_dimensions[2].height = 68
    hoja.row_dimensions[3].height = 45
    hist.column_dimensions["B"].width = 54
    hist.freeze_panes = "B7"
    for numero in (1, 3, 4, 5):
        hist.row_dimensions[numero].height = 45
    hist.auto_filter.ref = f"A6:D{max(6, hist.max_row)}"
    etapas.auto_filter.ref = "A1:K5"
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def pdf_estadistico(resumen, filtros):
    from reportlab.graphics.charts.piecharts import Pie
    from reportlab.graphics.shapes import Drawing, Rect, String
    from reportlab.lib import colors
    from reportlab.platypus import PageBreak, Spacer, Table, TableStyle
    from .pdf_reportes import _documento, _parrafo, _tabla

    salida, doc, estilos, contenido, pagina = _documento("Indicadores de selección", filtros, resumen["total"])
    def tabla(cabeceras, filas, proporciones):
        resultado = _tabla(cabeceras, filas, proporciones, doc.width, estilos)
        resultado.setStyle(TableStyle([("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
        return resultado
    contenido.append(tabla(["Recibidos", "Contratados", "Rechazados", "Pendientes"], [[resumen["total"], resumen["contratados"], resumen["rechazados"], resumen["en_curso"] + resumen["listos"]]], [1, 1, 1, 1]))
    contenido.append(Spacer(1, 10))
    contenido.append(_parrafo("Selección por llegada de la hoja de vida. Los estados corresponden al momento de generar el informe, no a decisiones tomadas durante esas fechas.", estilos["Nota"]))
    paleta = [colors.HexColor(c) for c in ("#2563EB", "#8B5CC8", "#138366", "#D44759")]

    def dona(items, titulo):
        dibujo = Drawing(doc.width / 2 - 14, 140)
        dibujo.add(String(0, 126, titulo, fontName="Helvetica-Bold", fontSize=11, fillColor=colors.HexColor("#182D43")))
        total = sum(item["cantidad"] for item in items)
        if total:
            pie = Pie()
            pie.x, pie.y, pie.width, pie.height = 0, 10, 100, 100
            pie.data = [item["cantidad"] for item in items]
            pie.labels = None
            pie.innerRadiusFraction = .65
            pie.slices.strokeWidth = 0
            for i, tono in enumerate(paleta):
                pie.slices[i].fillColor = tono
            dibujo.add(pie)
            dibujo.add(String(50, 55, str(total), textAnchor="middle", fontName="Helvetica-Bold", fontSize=18))
        else:
            dibujo.add(String(0, 94, "Sin datos para estos filtros", fontName="Helvetica", fontSize=10))
        for i, item in enumerate(items):
            nombre = "Seguridad y Salud" if item["etiqueta"] == "Seguridad y Salud en el Trabajo" else item["etiqueta"]
            dibujo.add(Rect(120, 100 - i * 25, 8, 8, fillColor=paleta[i % 4], strokeColor=None))
            dibujo.add(String(135, 100 - i * 25, f"{nombre}: {item['cantidad']}", fontName="Helvetica", fontSize=9))
        return dibujo

    contenido.append(Spacer(1, 14))
    contenido.append(Table([[dona(resumen["graficas"]["estados"], "Estado de los candidatos"), dona(resumen["graficas"]["rechazos"], "Rechazos por etapa")]], colWidths=[doc.width / 2] * 2, hAlign="LEFT"))
    historico = resumen.get("historico", {})
    filas = historico.get("filas", [])
    if filas or historico.get("error"):
        from math import ceil
        from reportlab.graphics.shapes import Line
        maximo = max(4, ceil(max((f[k] for f in filas for k in ("recibidos", "contratados", "rechazados")), default=0) / 4) * 4)
        for inicio in range(0, max(1, len(filas)), 12):
            bloque = filas[inicio:inicio + 12]
            contenido.append(PageBreak())
            contenido.append(_parrafo("Histórico de recibidos, contratados y rechazados", estilos["Heading2"]))
            contenido.append(_parrafo(f"Agrupación: {historico.get('agrupacion', 'mensual')}. {historico.get('desde', '-')} a {historico.get('hasta', '-')}. Recibidos por llegada HDV; resultados por fecha del cierre válido, aunque llegaron antes del rango.", estilos["Nota"]))
            if historico.get("error"):
                contenido.append(_parrafo(historico["error"], estilos["Nota"]))
                break
            dibujo = Drawing(doc.width, 160)
            izquierda, base, alto = 32, 30, 100
            ancho = doc.width - izquierda - 12
            for i in range(5):
                y = base + alto * i / 4
                dibujo.add(Line(izquierda, y, doc.width - 12, y, strokeColor=colors.HexColor("#DCE5EE"), strokeWidth=.5))
                dibujo.add(String(izquierda - 6, y - 3, str(int(maximo * i / 4)), textAnchor="end", fontName="Helvetica", fontSize=8))
            series = [("recibidos", "Recibidos", paleta[0]), ("contratados", "Contratados", paleta[2]), ("rechazados", "Rechazados", paleta[3])]
            paso = ancho / max(1, len(bloque))
            for i, fila in enumerate(bloque):
                centro = izquierda + paso * (i + .5)
                for j, (clave, _, tono) in enumerate(series):
                    valor = fila[clave]
                    x, altura = centro + (j - 1.5) * 12, valor / maximo * alto
                    dibujo.add(Rect(x, base, 10, altura, fillColor=tono, strokeColor=None))
                    if valor:
                        dibujo.add(String(x + 5, base + altura + 4, str(valor), textAnchor="middle", fontName="Helvetica", fontSize=8))
                dibujo.add(String(centro, 15, fila["etiqueta"].replace(" · ", " - "), textAnchor="middle", fontName="Helvetica", fontSize=7))
            for i, (_, nombre, tono) in enumerate(series):
                dibujo.add(Rect(izquierda + i * 140, 148, 8, 8, fillColor=tono, strokeColor=None))
                dibujo.add(String(izquierda + 14 + i * 140, 148, nombre, fontName="Helvetica", fontSize=8))
            contenido.append(dibujo)
            contenido.append(tabla(["Periodo", "Recibidos", "Contratados", "Rechazados"], [[f["etiqueta"], f["recibidos"], f["contratados"], f["rechazados"]] for f in bloque], [2, 1, 1, 1]))
            if historico.get("sin_fecha"):
                contenido.append(_parrafo(f"{historico['sin_fecha']} procesos finalizados sin fecha de cierre verificable; no se inventan fechas.", estilos["Nota"]))
    contenido.append(PageBreak())
    contenido.append(_parrafo("Tiempos y cumplimiento por etapa", estilos["Title"]))
    contenido.append(_parrafo("Solo etapas finalizadas; se excluyen cierres administrativos de ajuste. Horario hábil: 7:00 a. m. a 4:00 p. m., sin fines de semana ni festivos de Colombia.", estilos["Nota"]))
    contenido.append(Spacer(1, 12))
    def horas(valor):
        return f"{valor:g} h" if valor is not None else "Sin datos"
    contenido.append(_tabla(["Etapa", "Plazo", "Promedio", "Mediana", "Finalizadas", "A tiempo", "Fuera de plazo"], [[e["etapa"], f"{e['plazo_dias']} días", horas(e["promedio_horas"]), horas(e["mediana_horas"]), e["finalizadas"], e["a_tiempo"], e["retrasos_cerrados"]] for e in resumen["por_etapa"]], [3, 1, 1, 1, 1, 1, 1], doc.width, estilos))
    contenido.append(_parrafo("Vacantes con más candidatos", estilos["Heading2"]))
    contenido.append(_tabla(["Vacante", "Candidatos"], [[v["vacante"], v["cantidad"]] for v in resumen["vacantes"]], [5, 1], doc.width, estilos))
    contenido.append(Spacer(1, 12))
    contenido.append(_parrafo(f"Pendientes de atención: {resumen['vencidos']} etapas abiertas vencidas y {len(resumen['proximos_vencer'])} próximas a vencer. Las alertas cambian con el transcurso del tiempo.", estilos["Nota"]))
    doc.build(contenido, onFirstPage=pagina, onLaterPages=pagina)
    return salida.getvalue()
