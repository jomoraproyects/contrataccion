from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from reportlab.lib.styles import getSampleStyleSheet

from .models import Decision, Perfil, ProcesoSeleccion
from .reportes import _detalle_rechazos_pdf


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class ChecklistConsultaReportesTests(TestCase):
    def setUp(self):
        self.usuarios = {}
        for rol in (Perfil.Rol.VINCULACION, Perfil.Rol.PSICOLOGIA, Perfil.Rol.GERENTE):
            usuario = User.objects.create_user(f"reporte_{rol.lower()}", password="prueba")
            usuario.perfil.rol = rol
            usuario.perfil.save()
            self.usuarios[rol] = usuario
        self.proceso = ProcesoSeleccion.objects.create(
            nombre="María", apellidos="Verificación", cedula="88991234", celular="3001234567",
            fecha_llegada=timezone.localdate(), vacante="Operaciones", etapa_actual="PSICOLOGIA",
            creado_por=self.usuarios[Perfil.Rol.VINCULACION],
        )
        self.proceso.registrar_decision(
            self.usuarios[Perfil.Rol.PSICOLOGIA], Decision.Resultado.RECHAZADO,
            "El resultado de la prueba no cumple con el perfil requerido",
            {"entrevista": "CUMPLE", "prueba_psicotecnica": "NO_CUMPLE"},
        )

    def test_gerencia_ve_fallo_y_lista_completa_en_el_historial(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        respuesta = self.client.get(reverse("procesos:detalle", args=[self.proceso.pk]))
        self.assertContains(respuesta, "No cumplió:")
        self.assertContains(respuesta, "Prueba psicotécnica")
        self.assertContains(respuesta, "Ver lista de verificación")
        self.assertContains(respuesta, "Entrevista")
        self.assertContains(respuesta, "El resultado de la prueba no cumple con el perfil requerido")
        self.assertNotContains(respuesta, "data-stage-checklist")

    def test_consulta_cedula_y_registro_nuevo_muestran_requisito_fallido(self):
        self.client.force_login(self.usuarios[Perfil.Rol.VINCULACION])
        datos = self.client.get(reverse("procesos:consultar_candidato_cedula"), {"cedula": self.proceso.cedula}).json()
        self.assertEqual(datos["ultimo_rechazo"]["requisitos"], ["Prueba psicotécnica"])
        respuesta = self.client.post(reverse("procesos:crear"), {
            "nombre": "María", "apellidos": "Verificación", "cedula": self.proceso.cedula,
            "celular": self.proceso.celular, "fecha_llegada": timezone.localdate().isoformat(), "vacante": "Ventas",
        })
        self.assertContains(respuesta, "No cumplió: Prueba psicotécnica")
        self.assertEqual(ProcesoSeleccion.objects.count(), 1)

    def test_excel_y_pdf_incluyen_fallos_y_conservan_observacion_separada(self):
        self.client.force_login(self.usuarios[Perfil.Rol.GERENTE])
        excel = self.client.get(reverse("procesos:exportar_indicadores", args=["xlsx"]))
        hoja = load_workbook(BytesIO(excel.content), read_only=True)["Candidatos"]
        self.assertEqual(hoja["J1"].value, "Requisitos no cumplidos")
        self.assertEqual(hoja["J2"].value, "Prueba psicotécnica")
        self.assertEqual(hoja["H2"].value, "Psicología")
        self.assertIn("resultado de la prueba", hoja["I2"].value)
        traza = self.client.get(reverse("procesos:trazabilidad"), {"exportar": "xlsx"})
        hoja_traza = load_workbook(BytesIO(traza.content), read_only=True)["Candidatos"]
        self.assertEqual(hoja_traza["M2"].value, "Prueba psicotécnica")
        bloques = _detalle_rechazos_pdf([self.proceso], getSampleStyleSheet())
        self.assertTrue(any("Prueba psicotécnica" in getattr(bloque, "text", "") for bloque in bloques))
        pdf = self.client.get(reverse("procesos:exportar_indicadores", args=["pdf"]))
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))
